"""Replace report formula blocks with editable Word OMML, without changing values."""
import ast
import re
import sys
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT


def el(name, *children):
    e = OxmlElement('m:' + name)
    for c in children:
        e.append(c)
    return e


def run(value, roman=False):
    r = el('r')
    if roman:
        prop = el('rPr'); style = el('sty'); style.set(qn('m:val'), 'p'); prop.append(style); r.append(prop)
    t = el('t'); t.text = str(value); r.append(t)
    return r


def name(value):
    value = {'rho':'ρ','alpha':'α','beta':'β','pi':'π','lam':'λ','sumzeta':'Σζ','dz':'Δz','eta_p':'η_p'}.get(value,value)
    if '_' not in value:
        return [run(value, value in ['AOR','SOR','SLR','HRT','SRT','MLSS','MLVSS','DO','TN','TP','BOD','COD','RAS','IR'])]
    base, sub = value.split('_',1)
    if sub == 'max_h': sub = 'max,h'
    return [el('sSub',el('e',run(base, len(base)>1)),el('sub',run(sub,True)))]


def group(items):
    return el('d',el('e',*items))


def expr(node):
    if isinstance(node,ast.Name): return name(node.id)
    if isinstance(node,ast.Constant): return [run(node.value)]
    if isinstance(node,ast.UnaryOp): return [run('−'),*expr(node.operand)]
    if isinstance(node,ast.BinOp):
        a,b=expr(node.left),expr(node.right)
        if isinstance(node.op,ast.Div): return [el('f',el('num',*a),el('den',*b))]
        if isinstance(node.op,ast.Pow):
            base = [group(a)] if isinstance(node.left,ast.BinOp) else a
            return [el('sSup',el('e',*base),el('sup',*b))]
        if isinstance(node.op,ast.Mult):
            if isinstance(node.left,ast.BinOp) and isinstance(node.left.op,(ast.Add,ast.Sub)): a=[group(a)]
            if isinstance(node.right,ast.BinOp) and isinstance(node.right.op,(ast.Add,ast.Sub)): b=[group(b)]
            return [*a,run('·'),*b]
        return [*a,run('+' if isinstance(node.op,ast.Add) else '−'),*b]
    if isinstance(node,ast.Call) and node.func.id=='sqrt':
        pr=el('radPr');hide=el('degHide');hide.set(qn('m:val'),'1');pr.append(hide)
        return [el('rad',pr,el('deg'),el('e',*expr(node.args[0])))]
    if isinstance(node,ast.Call) and node.func.id=='sin':
        return [el('func',el('fName',run('sin',True)),el('e',*expr(node.args[0])))]
    if isinstance(node,ast.Call) and node.func.id=='deg':
        return [*expr(node.args[0]),run('°',True)]
    raise ValueError(ast.dump(node))


def parse(value): return expr(ast.parse(value,mode='eval').body)


FORMULAS = {
 '2-1':['Q_d=N*q/1000'], '2-2':['Q_max_h=Kz*Q_h'],
 '4-1':['v=Q*sqrt(sin(alpha))/(b*n*h)'], '4-2':['W_c=Q_d*w_c/1000'],
 '5-1':['V_min=Q_p*t_p','h_min=V_min/(L*B)'], '5-2':['V_cell=L*B*h'],
 '5-3':['H_req=dz+h_branch+h_main+h_add'], '5-4':['P_shaft=rho*g*Q_p*H/(1000*eta_p)'],
 '6-1':['v_H=Q/(B*h)'], '6-2':['t=V/(60*Q)','V=L*B*h'],
 '7-1':['q_s=Q_max_h/(n_c*L*B)'], '7-2':['t=L*B*h/(Q_max_h/n_c)'], '7-3':['v_H=Q_c/(B*h)'],
 '8-1':['L_MLSS=M_BOD/(V*X)'], '8-2':['L_MLVSS=M_BOD/(V*X_v)'],
 '9-1':['t_n=V_n/Q_h'],
 '10-1':['M_rem=Q_d*(S_0-S_e)/1000','O_d=k_O*M_rem'],
 '10-2':['AOR_design=(O_d/24)*Kz*K_m'],
 '10-3':['f=alpha*(beta*C_sm-C_L)*1.024**(T-20)/C_20','SOR=AOR_design/f'],
 '10-4':['G_20=SOR/(rho_O*E_A)','G_N=G_20*T_N/T_20'],
 '10-5':['n_d=G_adopt/g_d','n_series=n_d/4'],
 '11-1':['A_net=n_s*pi*(D_s**2-D_c**2)/4'], '11-2':['SLR_max=Q_eq*(1+R)*X/A_net'],
 '12-1':['A=pi*D**2/4','v=Q/A','h_f=lam*(L/D)*v**2/(2*g)','h_l=sumzeta*v**2/(2*g)'],
 '12-2':['h_f=L*(n_M*v/R_h**(2/3))**2'], '12-3':['h_w=(Q/(1.84*b_w))**(2/3)'],
 '15-1':['V_req=Q_max_h*t_c/60','t_max=60*V_total/Q_max_h'], '15-2':['y_c=(Q**2/(g*b**2))**(1/3)'],
}


def equation(items):
    result=[]
    for i,item in enumerate(items):
        if i: result.append(run('；　',True))
        left,right=item.split('=',1)
        result.extend([*parse(left),run('='),*parse(right)])
    return result


# Full expressions in calculation prose and tables also remain native equations.
SYMBOLS = sorted(set(re.findall(r'[A-Za-zα-ωΑ-Ω][A-Za-z0-9_]*',
    ' '.join(v for values in FORMULAS.values() for v in values))) |
    set('HRT SRT MLSS MLVSS DO TN TP BOD COD RAS IR AOR SOR SLR Q_max,h Q_d Q_h Q_s Q_c Q_p Q_in Q_out Q_eq H_s n_c n_s n_d n_M v_lim h_b h_0 β_g ξ α β ρ λ π η Δh Δp_d Δp_pipe Δp_m C_in C_out M_Cl D_Cl c_Cl V_sol V_store a_s G_a I_a V_h L_h B_t B_b A_i L_i B_i A_road A_used A_other V_j t_j L_w D_w q_w E_c P_in P_shaft η_m Σζ'.split()),key=len,reverse=True)


def normalize(value):
    value=value.replace('Q_max,h','Q_max_h').replace('−','-').replace('×','*').replace('²','**2').replace('³','**3').replace('^','**').replace('[','(').replace(']',')')
    value=re.sub(r'sin(\d+)°',r'sin(deg(\1))',value)
    value=value.replace('sinα','sin(alpha)').replace('√','sqrt')
    tokens=re.findall(r'\d+(?:\.\d+)?|[A-Za-zα-ωΑ-ΩΣ][A-Za-zα-ωΑ-ΩΣ0-9_]*|\*\*|.',value)
    expanded=[]
    for token in tokens:
        if re.match(r'^[A-Za-zα-ωΑ-ΩΣ]',token) and token not in ['sqrt','sin','deg']:
            remaining=token;pieces=[]
            while remaining:
                hit=next((s for s in SYMBOLS if remaining.startswith(s)),None)
                if hit is None: pieces=[token];break
                pieces.append(hit);remaining=remaining[len(hit):]
            expanded.extend(pieces)
        else: expanded.append(token)
    result=[]
    for i,token in enumerate(expanded):
        if i:
            prev=expanded[i-1]
            left=bool(re.match(r'^[\wα-ωΑ-ΩΣ]',prev)) or prev==')'
            right=bool(re.match(r'^[\wα-ωΑ-ΩΣ]',token)) or token=='('
            if left and right and not(prev in ['sqrt','sin','deg'] and token=='('):result.append('*')
        result.append(token)
    return ''.join(result)


INLINE = re.compile(r'[A-Za-zα-ωΑ-ΩΣ0-9_(\[][A-Za-zα-ωΑ-ΩΣ0-9_,.()+*/=<>^×−√°²³\[\]-]*')


def inline_math(p):
    if p._p.xpath('.//m:oMath') or p._p.xpath('.//w:drawing'):return 0
    value=p.text;replacements=[]
    if 'BOD₅/SS/TN' in value:return 0  # Slash-separated indicator list is not division.
    for match in INLINE.finditer(value):
        candidate=match[0].rstrip('.,')
        if '=' not in candidate or candidate.count('=')<1:continue
        if candidate.startswith('x=') and candidate.count('/')>1:continue
        try:
            nodes=[]
            for i,part in enumerate(candidate.split('=')):
                if i:nodes.append(run('='))
                nodes.extend(parse(normalize(part)))
            replacements.append((match.start(),match.start()+len(candidate),nodes))
        except (SyntaxError,ValueError,AttributeError):continue
    if not replacements:return 0
    p.clear();cursor=0
    for start,end,nodes in replacements:
        p.add_run(value[cursor:start]);p._p.append(el('oMath',*nodes));cursor=end
    p.add_run(value[cursor:]);return len(replacements)


def apply(path):
    d=Document(path)
    width=d.sections[0].page_width-d.sections[0].left_margin-d.sections[0].right_margin
    found=[]
    for p in d.paragraphs:
        match=re.search(r'（(\d+-\d+)）$',p.text)
        if not match or match[1] not in FORMULAS: continue
        key=match[1];found.append(key);p.clear()
        p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        fmt=p.paragraph_format;fmt.first_line_indent=Pt(0);fmt.left_indent=Pt(0)
        fmt.space_before=Pt(4);fmt.space_after=Pt(4);fmt.line_spacing=1.15
        fmt.tab_stops.clear_all();fmt.tab_stops.add_tab_stop(int(width/2),WD_TAB_ALIGNMENT.CENTER)
        fmt.tab_stops.add_tab_stop(width,WD_TAB_ALIGNMENT.RIGHT)
        p.add_run('\t')
        parts=FORMULAS[key]
        if key=='12-1':
            arr=el('eqArr',el('e',*equation(parts[:2])),el('e',*equation(parts[2:])))
            math=el('oMath',arr)
        else: math=el('oMath',*equation(parts))
        p._p.append(math);p.add_run('\t（'+key+'）')
    assert set(found)==set(FORMULAS),(found,set(FORMULAS)-set(found))
    # Typeset the symbols in the legend with the same native subscript objects.
    count_symbols=0
    for p in d.paragraphs:
        value=p.text
        if not value.startswith(('式中：','　　　')) or '——' not in value: continue
        prefix='式中：' if value.startswith('式中：') else '　　　'
        variable,rest=value[len(prefix):].split('——',1)
        if not re.fullmatch(r'[A-Za-zΑ-ωΔΣρλζηαβπ_0-9、,]+',variable): continue
        p.clear();p.add_run(prefix)
        symbols=re.split('([、,])',variable.replace('_max,h','_max_h'))
        math=el('oMath')
        for token in symbols:
            if token in ['、',',']: math.append(run(token,True))
            elif token:
                for e in name(token): math.append(e)
        p._p.append(math);p.add_run('——'+rest);count_symbols+=1
    mathpr=d.settings._element.find(qn('m:mathPr'))
    if mathpr is None: mathpr=el('mathPr');d.settings._element.append(mathpr)
    font=mathpr.find(qn('m:mathFont'))
    if font is None: font=el('mathFont');mathpr.append(font)
    font.set(qn('m:val'),'Cambria Math')
    inline_count=sum(inline_math(p) for p in d.paragraphs)
    for table in d.tables:
        if table.cell(0,0).text=='工况' and '次氯酸钠' in table.cell(0,1).text:
            for row in table.rows[:-1]:
                for cell in row.cells:
                    for p in cell.paragraphs:p.paragraph_format.keep_with_next=True
        seen=set()
        for row in table.rows:
            for cell in row.cells:
                if cell._tc in seen:continue
                seen.add(cell._tc)
                inline_count+=sum(inline_math(p) for p in cell.paragraphs)
                for p in cell.paragraphs:
                    for r in p.runs:r.font.size=Pt(9)
                    for r in p._p.xpath('.//m:r'):
                        pr=OxmlElement('w:rPr');sz=OxmlElement('w:sz');sz.set(qn('w:val'),'18');pr.append(sz)
                        r.insert(1 if len(r) and r[0].tag==qn('m:rPr') else 0,pr)
    d.save(path)
    print(f'Native Word formula blocks: {len(found)}; native legend symbols: {count_symbols}; inline calculations: {inline_count}')


if __name__=='__main__': apply(Path(sys.argv[1]))
