"""Append the authoritative teacher-review chapter to the staged DOCX report.

Run once against the repository version.  Existing formulas and historical
sections remain for traceability; chapter 16 explicitly supersedes conflicting
values and is generated from the reviewed calculation results.
"""
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_BREAK, WD_PARAGRAPH_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "计算成果" / "污水处理厂设计计算说明书_阶段核验稿.docx"


def iter_paragraphs(document):
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def replace_runs(paragraph, old, new):
    for run in paragraph.runs:
        if old in run.text:
            run.text = run.text.replace(old, new)


def set_cell(cell, text, bold=False):
    cell.text = str(text)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    if bold and paragraph.runs:
        paragraph.runs[0].bold = True


def add_table(document, headers, rows, widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        set_cell(table.rows[0].cells[index], header, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell(cells[index], value)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Cm(width)
    document.add_paragraph()
    return table


def add_body(document, text, bold_prefix=None):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Cm(0.74)
    paragraph.paragraph_format.line_spacing = 1.5
    if bold_prefix and text.startswith(bold_prefix):
        paragraph.add_run(bold_prefix).bold = True
        paragraph.add_run(text[len(bold_prefix):])
    else:
        paragraph.add_run(text)
    return paragraph


def main():
    document = Document(DOCX)
    already_has_chapter = any("教师审阅修订汇总" in paragraph.text for paragraph in document.paragraphs)

    document.core_properties.title = "污水处理厂设计计算说明书（教师审阅修订稿）"
    document.core_properties.subject = "2026-09-10教师录音意见修订"

    replacements = {
        "污水处理厂设计计算说明书（阶段核验稿）": "污水处理厂设计计算说明书（教师审阅修订稿）",
        "4座D44": "6座D36",
        "4 座 D44": "6 座 D36",
        "3840个": "6000个",
        "3840 个": "6000 个",
        "960/系列": "1500/系列",
        "960 个/系列": "1500 个/系列",
        "7.8125 m³/h/盘": "5.0 Nm³/h/盘",
        "9 m/90 kW": "13.0 m/132 kW",
        "9.0 m、90 kW": "13.0 m、132 kW",
    }
    for paragraph in iter_paragraphs(document):
        for old, new in replacements.items():
            replace_runs(paragraph, old, new)
        if paragraph.text.strip() == "阶段核验稿":
            paragraph.text = "教师审阅修订稿"
    for section in document.sections:
        for paragraph in list(section.header.paragraphs) + list(section.footer.paragraphs):
            if "阶段核验稿" in paragraph.text:
                paragraph.text = paragraph.text.replace("阶段核验稿", "教师审阅修订稿")
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip() in {"9 m，90 kW", "9 m, 90 kW"}:
                    set_cell(cell, "13.0 m，132 kW")

            # Keep the front matter and the closing summary usable without
            # forcing readers to discover the superseding values in chapter 16.
            cells = row.cells
            label = cells[0].text.strip() if cells else ""
            if label == "提升泵" and len(cells) == 4:
                set_cell(cells[2], "13.0 m、132 kW")
                set_cell(cells[3], "第13项教师审阅复算")
            elif label == "提升泵" and len(cells) == 3:
                set_cell(cells[1], "示意图未完整表达；9 m/90 kW")
            elif label == "表面负荷" and len(cells) == 4:
                set_cell(cells[1], "1.523")
                set_cell(cells[2], "1.979")
                set_cell(cells[3], "按扣除D4.5中心井净面积，满足")
            elif label == "停留时间" and len(cells) == 4:
                set_cell(cells[1], "1.970 h")
                set_cell(cells[2], "1.516 h")
                set_cell(cells[3], "按3.00 m沉淀区，满足")
            elif label == "D/h" and len(cells) == 4:
                set_cell(cells[2], "10.33")
                set_cell(cells[3], "6～12，满足")
            elif label == "曝气器" and len(cells) == 3:
                set_cell(cells[1], "6000个")
                set_cell(cells[2], "Sanitaire Silver Series II 229 mm；1500个/系列、5.0 Nm³/h·个")
            elif label == "鼓风机" and len(cells) == 3:
                set_cell(cells[1], "6台")
                set_cell(cells[2], "AERZEN AT200-0.8 G5plus；6000 m³/h、65 kPa、165 kW，5用1备")
            elif label == "单泵出水支管" and len(cells) == 6:
                for index, value in enumerate(("0.6250", "DN900", "12", "0.982", "0.225"), start=1):
                    set_cell(cells[index], value)
            elif label == "初沉配水—单池" and len(cells) == 6:
                for index, value in enumerate(("0.40625", "DN900", "48", "0.639", "0.064"), start=1):
                    set_cell(cells[index], value)
            elif label == "初沉—单生物系列" and len(cells) == 6:
                for index, value in enumerate(("0.40625", "DN900", "15", "0.639", "0.038"), start=1):
                    set_cell(cells[index], value)
            elif label == "单生物—二沉配水" and len(cells) == 6:
                for index, value in enumerate(("0.71094", "DN1000", "59", "0.905", "0.154"), start=1):
                    set_cell(cells[index], value)
            elif label == "二沉配水—单池" and len(cells) == 6:
                for index, value in enumerate(("0.47396", "DN1000", "106", "0.603", "0.076"), start=1):
                    set_cell(cells[index], value)
            elif label == "单二沉出水" and len(cells) == 6:
                for index, value in enumerate(("0.27083", "DN800", "25", "0.539", "0.039"), start=1):
                    set_cell(cells[index], value)
            elif label == "二沉出水干管A/B" and len(cells) == 6:
                for index, value in enumerate(("0.8125", "DN1200", "310/120", "0.718", "0.215/0.118"), start=1):
                    set_cell(cells[index], value)
            elif label == "初沉配水井" and len(cells) == 3:
                set_cell(cells[2], "峰值停留44.3 s；2.00 m堰上水头0.230 m")
            elif label == "二沉配水井" and len(cells) == 3:
                set_cell(cells[1], "每线16×10 m占地包络，3条2.00 m堰")
                set_cell(cells[2], "单池峰值0.47396 m³/s；堰上水头0.255 m")
            elif label == "RAS" and len(cells) == 3:
                set_cell(cells[2], "最不利动态损失0.6525 m；设计扬程不小于2.35 m")

            elevation_rows = {
                "入厂D1200": ("平均24.041、峰值24.200", "—", "23.30管底", "27.30地面"),
                "湿井": ("22.60～23.80", "泵前额度0.30", "20.40；局部泵坑19.80", "27.30"),
                "细格栅": ("34.30/34.15", "0.150", "33.35", "34.85"),
                "曝气沉砂": ("34.00/33.75", "0.250", "31.25；砂斗另低1.50", "34.50"),
                "初沉配水井": ("33.50/33.15", "0.350", "32.00/31.65", "34.00"),
                "初沉池": ("33.00", "0.150", "29.50；斗底25.8375", "33.50"),
                "A²/O": ("32.60/32.30", "0.300", "27.30", "32.80"),
                "二沉配水井": ("31.95/31.55", "0.400", "29.95/29.55", "32.45"),
                "二沉池": ("31.40", "0.150", "26.40；斗底22.15", "31.90"),
                "消毒/控制井": ("30.90/30.40", "0.500", "暂27.90", "31.70"),
            }
            if label in elevation_rows and len(cells) == 5:
                for index, value in enumerate(elevation_rows[label], start=1):
                    set_cell(cells[index], value)

            if label == "粗/细格栅槽" and len(cells) == 4:
                set_cell(cells[1], "各2，正常2格并联")
                set_cell(cells[3], "两格同时承担单线峰值")
            elif label == "初沉池" and len(cells) == 4:
                set_cell(cells[1], "4")
                set_cell(cells[2], "净D31，中心井D4.5，沉淀深3.00 m")
                set_cell(cells[3], "峰值净表面负荷1.979")
            elif label == "提升潜污泵" and len(cells) == 4:
                set_cell(cells[1], "2250 m³/h、13.0 m、132 kW")
                set_cell(cells[3], "4台3用1备；曲线待厂家核")
            elif label == "生物鼓风机" and len(cells) == 4:
                set_cell(cells[1], "AERZEN AT200-0.8 G5plus；6000 m³/h、65 kPa、165 kW")
                set_cell(cells[3], "设计点须由厂家确认")
            elif label == "微孔盘曝气器" and len(cells) == 4:
                set_cell(cells[1], "Sanitaire Silver Series II 229 mm；5.0 Nm³/h/盘")
                set_cell(cells[2], "6000个；1500个/系列")
                set_cell(cells[3], "官方范围0.8～6.8 Nm³/h")
            elif label == "初沉刮泥机" and len(cells) == 4:
                set_cell(cells[1], "适配D31")
            elif (
                label == "初沉池"
                and len(cells) == 3
                and table.rows[0].cells[0].text.strip() == "校核对象"
            ):
                set_cell(cells[1], "扣中心井后峰值负荷1.979、停留1.516 h")
                set_cell(cells[2], "满足课程采用范围")
            elif (
                label == "初沉池"
                and len(cells) == 3
                and table.rows[0].cells[0].text.strip() == "审阅问题"
            ):
                set_cell(cells[1], "4座D28；峰值2.375")
                set_cell(cells[2], "4座净D31；扣D4.5中心井后峰值1.979")
            elif label == "V0.2" and len(cells) == 4:
                set_cell(cells[1], "2026-09-10")
                set_cell(cells[2], "按教师录音意见修订第3—14项及复算附件")
                set_cell(cells[3], "教师审阅修订稿")

    # Prevent a table header or the newly added revision headings from being
    # stranded at the bottom of a page after the longer reviewed values wrap.
    for table in document.tables:
        for paragraph in table.rows[0].cells[0].paragraphs:
            paragraph.paragraph_format.keep_with_next = True
        for cell in table.rows[0].cells[1:]:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True
    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith("16.") or paragraph.text.strip().startswith("表14-3"):
            paragraph.paragraph_format.keep_with_next = True

    if already_has_chapter:
        document.save(DOCX)
        print(DOCX)
        return

    document.add_page_break()
    heading = document.add_paragraph("16 教师审阅修订汇总", style="Heading 1")
    heading.paragraph_format.keep_with_next = True
    intro = document.add_paragraph()
    intro.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY
    intro.add_run("修订效力：").bold = True
    intro.add_run(
        "本章根据2026年9月10日教师录音意见编制。凡前文阶段计算值与本章冲突，"
        "以本章及仓库同日修订的第3—14项Markdown正文和复算附件为准。"
    )

    document.add_paragraph("16.1 主要修改及统一值", style="Heading 2")
    add_table(
        document,
        ["审阅问题", "原阶段值/表达", "教师审阅后统一值"],
        [
            ("格栅运行方式", "每线1用1备、每级4格", "每级2格、正常2格同时运行；检修期限流/跨线"),
            ("提升泵", "示意图未完整表达；9 m/90 kW", "4台3用1备；2250 m³/h×13.0 m；132 kW"),
            ("初沉池", "4座D28；峰值2.375", "4座净D31；扣D4.5中心井后峰值1.979"),
            ("A²/O系列", "两线各2系列，表述易混", "全厂4个独立系列；每系列54×45×5 m"),
            ("脱氮/补碳", "仅列碳源缺口", "80%甲醇最大设备能力5359.9 kg/d，按NO₃-N/TN反馈"),
            ("化学除磷", "仅预留接口", "40% FeCl₃ 5.922 m³/d；含20%余量取7.11 m³/d"),
            ("曝气设备", "3840盘、7.81 m³/h·盘", "Sanitaire Silver Series II 229 mm；6000盘、5.0 Nm³/h·盘"),
            ("鼓风机", "仅给6000 m³/h、65 kPa", "AERZEN AT200-0.8 G5plus；6台5用1备"),
            ("二沉总图", "4座D44预留", "6座净D36；结构外径D36.8；占地包络D40"),
            ("高程", "A²/O池底EL23.00", "A²/O池底EL27.30；主体水面EL32.30"),
        ],
        [3.2, 4.4, 8.2],
    )

    document.add_paragraph("16.2 初沉池及污泥斗", style="Heading 2")
    add_body(document, "全厂4座中心进水辐流初沉池。净D31 m、中心井D4.50 m，扣井后单池净沉淀面积738.86 m²、全厂2955.45 m²；平均/峰值表面负荷1.523/1.979 m³/(m²·h)，峰值停留1.516 h。")
    add_body(document, "每池污泥斗上口D4.50 m、下口D1.00 m、深3.00 m，斗壁角59.74°，容积20.22 m³；按4 h储泥所需15.625 m³，容积满足。初沉水面EL33.00 m，周边工艺内底EL29.50 m，斗底EL25.8375 m。")

    document.add_paragraph("16.3 A²/O分区、氮磷和药剂", style="Heading 2")
    add_body(document, "每系列分6条54×7.5 m廊道；等效流程长度按厌氧45 m、缺氧90 m、好氧189 m，总长324 m=6×54 m。对应容积1687.5/3375/7087.5 m³和HRT 1.5/3.0/6.3 h。")
    add_body(document, "TN需去除2640 kgN/d，其中同化391.094 kgN/d、反硝化2248.906 kgN/d。因COD组分缺失，补碳设备按污水自身可利用碳源贡献为零的上限计算：6431.87 kgCOD/d，折合80%甲醇5359.9 kg/d。该值是设备上限，实际必须反馈下调。")
    add_body(document, "TP需去除447.0 kgP/d；按4%生物带磷后化学缺口316.64 kgP/d。Fe/P摩尔比2.0时，纯FeCl₃ 3316.3 kg/d；40%溶液、密度1.40 kg/L为5.922 m³/d，含20%余量取7.11 m³/d。化学固体约2634 kg/d，使二沉峰值固体负荷约145.26 kg/(m²·d)，仍低于150。")

    document.add_paragraph("16.4 曝气器、鼓风机与依据", style="Heading 2")
    add_body(document, "微孔曝气器采用Xylem Sanitaire Silver Series II标准型229 mm作为课程候选；官方风量范围0.8—6.8 Nm³/h，本设计取5.0 Nm³/h，6000盘、每系列1500盘。总风量仍为30000 m³/h。")
    add_body(document, "鼓风机候选为AERZEN AT200-0.8 G5plus。官方型号包络8400 m³/h、80 kPa、165 kW；项目设计点6000 m³/h、65 kPa，共6台、5用1备。设计点及当地修正必须由厂家确认。")
    source = document.add_paragraph()
    source.paragraph_format.left_indent = Cm(0.74)
    source.add_run("产品资料：").bold = True
    source.add_run("Xylem Sanitaire Silver Series II 229 mm官方数据表；AERZEN Turbo Blower Generation 5plus官方型号页。")

    document.add_paragraph("16.5 总图、管线和高程", style="Heading 2")
    add_body(document, "总图已统一为4座D31初沉、4个A²/O系列、6座D36二沉和2座二沉配水井。初沉D34、二沉D40为含操作带的方形占地包络，不是池壁厚度。所有构筑物包络不重叠、不压道路并位于400×300 m厂界内。")
    add_body(document, "单泵支管采用DN900，流速0.982 m/s。12 m管长分解为湿井水平4 m、竖向7 m、设备中心线余量1 m；其余59 m、106 m、310/120 m均来自第12项坐标折线。题设固定进厂DN1200峰值流速1.786 m/s作为设计边界保留。")
    add_table(
        document,
        ["节点", "教师审阅后标高EL/m"],
        [
            ("湿井", "运行22.60—23.80；池底20.40"),
            ("细格栅", "栅前34.30；栅后34.15"),
            ("初沉", "水面33.00；周边底29.50；斗底25.8375"),
            ("A²/O", "入口32.60；主体32.30；池底27.30"),
            ("二沉", "水面31.40；周边底26.40；斗底22.15"),
            ("消毒入口/控制井", "30.90/30.40"),
        ],
        [5.0, 10.0],
    )
    add_body(document, "提升泵所需扬程12.858 m，选13.0 m；单泵输入功率102.32 kW，含10%余量112.55 kW，电机取132 kW。洪水位25.00 m、厂外DN1500暂按100 m时重力出厂成立。")

    document.add_paragraph("16.6 标准全称与复核边界", style="Heading 2")
    add_table(
        document,
        ["编号", "资料全称"],
        [
            ("GB 50014—2021", "《室外排水设计标准》"),
            ("GB 18918—2002", "《城镇污水处理厂污染物排放标准》"),
            ("HJ 576—2010", "《厌氧—缺氧—好氧活性污泥法污水处理工程技术规范》"),
            ("课程资料", "《给水排水设计手册 第1册 常用资料》《给水排水设计手册 第5册 城镇排水》《环境工程设计教案》《设计任务书6》"),
        ],
        [4.0, 11.0],
    )
    add_body(document, "23项数值工况中17项通过、6项为检修或Kz=1.5对照限制。仍须以实测COD组分、TKN/氨氮、碱度、最低温、烧杯试验、厂家曲线、实际组号及厂外管线测量完成正式设计闭合；课程计算不等同于施工图或实测达标证明。", bold_prefix="23项数值工况")

    # Ask Word/LibreOffice to refresh fields (TOC/page numbers) on open.
    settings = document.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")

    document.save(DOCX)
    print(DOCX)


if __name__ == "__main__":
    main()
