#!/usr/bin/env python3
"""Build the batch-8 submission DOCX and the 57-item final review workbook.

The document is rebuilt from the batch-7 illustrated work copy so that its
styles and frozen figures remain the visual baseline, while all narrative,
equations and tables are regenerated from the batch-2 through batch-8 frozen
interfaces.  This avoids leaving stale equations in hidden OOXML runs.
"""

from __future__ import annotations

import argparse
import json
import math
import re
import shutil
import zipfile
from copy import deepcopy
from datetime import date
from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


ROOT = Path(__file__).resolve().parents[1]
PROJECT = ROOT / "第二次修改协作项目_2026-09-13"
SOURCE_DOCX = (
    PROJECT
    / "02_成果回收"
    / "第09组_位图整图重绘"
    / "某污水处理厂设计计算说明书_第二次修改_整图重绘工作稿.docx"
)
FIGURE_DIR = PROJECT / "02_成果回收" / "第09组_位图整图重绘" / "导出PNG"
SOURCE_XLSX = PROJECT / "00_公共文件" / "修改与待修改_20260913.xlsx"
OUTPUT_DIR = PROJECT / "02_成果回收" / "第10组_全文统稿与终检"
OUTPUT_DOCX = OUTPUT_DIR / "某污水处理厂设计计算说明书_第二次修改稿.docx"
OUTPUT_XLSX = OUTPUT_DIR / "修改与待修改_20260914_终检版.xlsx"
OUTPUT_JSON = OUTPUT_DIR / "第10组_一致性核验结果.json"


FIGURES = {
    "3-1": "图3-1_水线工艺流程及回流排泥接口.png",
    "5-1": "图5-1_两格进水井与六泵位计算图.png",
    "6-1": "图6-1_曝气沉砂池计算图.png",
    "7-1": "图7-1_平流式初次沉淀池平剖计算图.png",
    "8-1": "图8-1_A2O单系列分区与回流泵接口图.png",
    "9-1": "图9-1_投药与固体闭合边界图.png",
    "10-1": "图10-1_曝气与鼓风系统计算图.png",
    "11-1": "图11-1_二次沉淀池平剖计算图.png",
    "12-1": "图12-1_主要连接管渠与配水接口图.png",
    "13-1": "图13-1_厂区总平面与管线图.png",
    "14-1": "图14-1_全厂水线高程控制图.png",
    "15-1": "图15-1_接触消毒与计量终端图.png",
}


def set_run_font(run, size: float = 10.5, bold: bool | None = None, name: str = "宋体") -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(0)

    body = doc.styles["Body Text"]
    body.font.name = "宋体"
    body._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    body.font.size = Pt(10.5)
    body.paragraph_format.first_line_indent = Pt(21)
    body.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    body.paragraph_format.space_after = Pt(0)
    body.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    h1 = doc.styles["Heading 1"]
    h1.font.name = "宋体"
    h1._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    h1.font.size = Pt(16)  # 三号
    h1.font.bold = True
    h1.font.color.rgb = RGBColor(0, 0, 0)
    h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h1.paragraph_format.first_line_indent = Pt(0)
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(8)
    h1.paragraph_format.keep_with_next = True

    h2 = doc.styles["Heading 2"]
    h2.font.name = "宋体"
    h2._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    h2.font.size = Pt(14)  # 四号
    h2.font.bold = True
    h2.font.color.rgb = RGBColor(0, 0, 0)
    h2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h2.paragraph_format.first_line_indent = Pt(0)
    h2.paragraph_format.space_before = Pt(8)
    h2.paragraph_format.space_after = Pt(5)
    h2.paragraph_format.keep_with_next = True

    h3 = doc.styles["Heading 3"]
    h3.font.name = "宋体"
    h3._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    h3.font.size = Pt(12)
    h3.font.bold = True
    h3.font.color.rgb = RGBColor(0, 0, 0)
    h3.paragraph_format.space_before = Pt(6)
    h3.paragraph_format.space_after = Pt(4)
    h3.paragraph_format.keep_with_next = True

    if "Title" in doc.styles:
        title = doc.styles["Title"]
        title.font.name = "宋体"
        title._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        title.font.size = Pt(22)
        title.font.bold = True
        title.font.color.rgb = RGBColor(0, 0, 0)


def clear_document_body(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def set_update_fields(doc: Document) -> None:
    settings = doc.settings._element
    node = settings.find(qn("w:updateFields"))
    if node is None:
        node = OxmlElement("w:updateFields")
        settings.append(node)
    node.set(qn("w:val"), "true")


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, value, end):
        run._r.append(node)
    set_run_font(run, 9)


def add_toc_field(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-2" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "打开文档后更新域即可显示目录"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, placeholder, end):
        run._r.append(node)
    set_run_font(run, 10.5)


def add_heading(doc: Document, text: str, level: int, page_break: bool = False):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    if page_break:
        # Put the break on the heading itself. A separate page-break paragraph
        # can be pushed onto a new page when the preceding table fills a page,
        # producing an unintended blank page in Microsoft Word.
        p.paragraph_format.page_break_before = True
    for run in p.runs:
        set_run_font(run, 16 if level == 1 else 14 if level == 2 else 12, True)
    return p


def add_body(doc: Document, text: str, indent: bool = True, bold_lead: str | None = None):
    p = doc.add_paragraph(style="Body Text")
    if not indent:
        p.paragraph_format.first_line_indent = Pt(0)
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_run_font(lead, 10.5, True)
        rest = p.add_run(text[len(bold_lead):])
        set_run_font(rest)
    else:
        run = p.add_run(text)
        set_run_font(run)
    return p


def add_formula(doc: Document, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_together = True
    run = p.add_run(text)
    set_run_font(run, 10.5, False, "Times New Roman")
    return p


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=70, start=80, bottom=70, end=80) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    flag = OxmlElement("w:tblHeader")
    flag.set(qn("w:val"), "true")
    tr_pr.append(flag)


def add_table(doc: Document, caption: str, headers: Sequence[str], rows: Iterable[Sequence[object]], widths=None):
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.keep_with_next = True
    run = cp.add_run(caption)
    set_run_font(run, 10.5, False)

    rows = [list(r) for r in rows]
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_repeat_table_header(table.rows[0])
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = str(header)
        shade_cell(cell, "D9EAF7")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            for r in p.runs:
                set_run_font(r, 9, True)
    for row_data in rows:
        row = table.add_row()
        for i, value in enumerate(row_data):
            cell = row.cells[i]
            cell.text = str(value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    set_run_font(r, 9)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_figure(doc: Document, key: str, caption: str) -> None:
    path = FIGURE_DIR / FIGURES[key]
    if not path.is_file():
        raise FileNotFoundError(path)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_together = True
    p.add_run().add_picture(str(path), width=Inches(6.45))
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.keep_with_next = True
    run = cp.add_run(caption)
    set_run_font(run, 10.5)


def add_cover(doc: Document) -> None:
    for _ in range(2):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("中国农业大学")
    set_run_font(r, 22, True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("环境工程课程设计")
    set_run_font(r, 22, True)
    for _ in range(2):
        doc.add_paragraph()
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("某污水处理厂设计计算说明书")
    set_run_font(r, 22, True)
    doc.add_paragraph()
    table = doc.add_table(rows=9, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    data = [
        ("学年学期", "2026—2027学年夏季学期"),
        ("学院", "资源与环境学院"),
        ("专业", "环境工程"),
        ("班级", "________________"),
        ("组号", "________________"),
        ("成员", "________________"),
        ("组内分工", "________________"),
        ("指导教师", "________________"),
        ("完成日期", "2026年9月"),
    ]
    for row, values in zip(table.rows, data):
        for cell, value in zip(row.cells, values):
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, 120, 120, 120, 120)
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    set_run_font(run, 12)
    doc.add_page_break()


def add_front_matter(doc: Document) -> None:
    add_heading(doc, "摘要", 1)
    add_body(
        doc,
        "本设计面向江苏省西南部某60万人城市污水处理厂。按人均排水量180 L/(人·d)计算平均日设计水量108000 m³/d（4500 m³/h），按GB 50014—2021取总变化系数Kz=1.50，最高时流量为6750 m³/h。处理目标采用课程任务书规定的一级B限值，并列明其与江苏省现行DB32/4440—2022 A标准之间的差异。",
    )
    add_body(
        doc,
        "污水处理水线采用粗格栅（地下）—两格进水井与提升泵—细格栅—曝气沉砂池—平流初沉池—A²/O生物池—二沉配水井—辐流二沉池—折流接触池—长喉计量渠—东厂界控制井。提升泵共6台、4用2备，单泵1687.5 m³/h、14 m、90 kW；A²/O设4个工艺系列，总有效容积63000 m³；回流污泥泵和内回流泵均为6台、4用2备。",
    )
    add_body(
        doc,
        "需氧量按HJ 576—2010的物料衡算方法计算，设计AOR为2049.209 kgO₂/h；标准空气量采用48000 Nm³/h，配置9600只229 mm盘式微孔曝气器和6台高速离心鼓风机（4用2备）。二沉采用8座净径37 m辐流池；消毒采用次氯酸钠，2格接触池总有效容积3456 m³。水力高程自细格栅前EL34.80 m降至计量后EL29.30 m，落差5.50 m；延伸至东厂界EL28.75 m后总落差6.05 m。",
    )
    p = doc.add_paragraph()
    r = p.add_run("关键词：城市污水；A²/O；脱氮除磷；曝气系统；水力高程；课程设计")
    set_run_font(r, 10.5)
    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("目录")
    set_run_font(r, 16, True)
    add_toc_field(doc)
def add_chapter_1(doc: Document) -> None:
    add_heading(doc, "1 前言", 1, True)
    add_heading(doc, "1.1 设计任务与成果边界", 2)
    add_body(doc, "任务书要求完成城市二级污水处理厂水线工艺设计，包括主要构筑物计算、工艺草图、总平面及高程关系。厂区面积12 hm²，厂坪标高EL27.30 m，北侧进水，进厂管为D1200、管底埋深4.00 m。本说明书按水线顺序给出设计输入、公式、代入、采用值与校核。")
    add_body(doc, "任务书未要求完整污泥处理线。本设计为使水线计算闭合，计算初沉污泥、剩余污泥、化学固体、回流及排泥边界；结果止于排泥接口，不代表已完成污泥浓缩、消化、脱水和处置系统设计。图件为课程设计计算图，施工图阶段仍需进行结构、设备、消防、除臭和管线综合设计。")
    add_heading(doc, "1.2 设计依据及适用范围", 2)
    add_table(doc, "表1-1 设计依据及用途", ["类别", "依据", "用途"], [
        ["课程资料", "设计任务书6(2)、环境工程设计教案", "规模、水质、厂址和计算书编制要求"],
        ["强制规范", "GB 55027—2022《城乡排水工程项目规范》", "排水工程安全、检修、消毒和计量底线"],
        ["设计标准", "GB 50014—2021《室外排水设计标准》", "流量、泵站、沉砂、沉淀、生物处理及水力设计"],
        ["排放标准", "GB 18918—2002及2025修改单", "任务书一级B目标的标准来源"],
        ["地方标准", "DB32/4440—2022", "江苏新建项目实际排放边界"],
        ["行业标准", "HJ 576—2010", "A²/O、回流、供氧和药剂系统设计"],
        ["设计手册", "《给水排水设计手册》第1册、第5册", "水力参数和构筑物资料复核"],
    ])
    add_body(doc, "依据采用顺序为全文强制通用规范、现行专业设计标准和适用排放标准。GB 55027—2022与其他工程建设标准不一致时优先执行；其公告只废止GB 50014—2021中列明的原强制条文，并非废止整本GB 50014—2021。")


def add_chapter_2(doc: Document) -> None:
    add_heading(doc, "2 设计水量、水质与处理程度", 1, True)
    add_heading(doc, "2.1 设计流量", 2)
    add_body(doc, "设计人口N=600000人，人均排水量q=180 L/(人·d)。平均日流量和最高时流量计算如下：")
    add_formula(doc, "Qavg=Nq=600000×180/1000=108000 m³/d=4500 m³/h=1.250 m³/s")
    add_formula(doc, "Qmax=KzQavg,h=1.50×4500=6750 m³/h=1.875 m³/s")
    add_table(doc, "表2-1 全厂流量口径", ["工况", "全厂流量", "单水线流量", "用途"], [
        ["平均日", "108000 m³/d；4500 m³/h", "2250 m³/h", "日负荷、HRT和物料衡算"],
        ["最高时", "6750 m³/h；1.875 m³/s", "3375 m³/h；0.9375 m³/s", "构筑物与连接管渠主算"],
        ["提升泵组合", "4×1687.5=6750 m³/h", "2×1687.5=3375 m³/h", "两条水线各2台工作泵"],
    ])
    add_heading(doc, "2.2 设计水质与出水目标", 2)
    add_table(doc, "表2-2 进水负荷及任务书出水目标", ["指标", "进水负荷 kg/d", "进水浓度 mg/L", "任务书目标 mg/L", "最低去除率"], [
        ["BOD₅", "18000", "166.67", "20", "88.0%"],
        ["SS", "30000", "277.78", "20", "92.8%"],
        ["TN", "4800", "44.44", "20", "55.0%"],
        ["TP", "600", "5.56", "1.0", "82.0%"],
        ["NH₄⁺-N", "未给定", "未给定", "8.0", "以TN负荷作容量上界"],
    ])
    add_table(doc, "表2-3 任务书目标与江苏A标准差异", ["指标", "任务书目标 mg/L", "DB32/4440—2022 A标准日均限值 mg/L", "结论"], [
        ["CODCr", "任务书未给出", "30", "缺进水COD，不能证明实际达标"],
        ["BOD₅", "20", "10", "课程目标较宽"],
        ["SS", "20", "10", "课程目标较宽"],
        ["TN", "20", "10（12）", "冬季括号值仍严于任务书"],
        ["NH₄⁺-N", "8", "1.5（3）", "课程目标较宽"],
        ["TP", "1.0", "0.3", "课程目标较宽"],
    ])
    add_body(doc, "本说明书保留任务书目标作为课程计算边界。按江苏新建、规模108000 m³/d的一般区域项目判断，实际工程应执行DB32/4440—2022 A标准，并需补充COD、TKN、氨氮、碱度、可利用碳源、最低连续水温和消毒试验后重新核定深度处理能力。")


def add_chapter_3(doc: Document) -> None:
    add_heading(doc, "3 工艺方案与总体流程", 1, True)
    add_heading(doc, "3.1 工艺比选与确定", 2)
    add_table(doc, "表3-1 生物处理方案比较", ["方案", "主要特点", "本项目适用性"], [
        ["A²/O", "厌氧、缺氧、好氧同池串联，可同步脱氮除磷", "流程成熟；适合本项目，碳源不足时可补甲醇"],
        ["氧化沟", "抗冲击能力强，构筑物占地较大", "12 hm²用地内布置紧张，分区控制不如A²/O直接"],
        ["SBR", "集反应与沉淀于一体，周期运行", "10.8万m³/d规模下池组、控制和连续配水较复杂"],
    ])
    add_body(doc, "综合水质、规模、占地和运行管理，本设计采用A²/O。初沉池降低SS和部分BOD₅；A²/O完成有机物去除、硝化反硝化与生物除磷；碳源不足时投加甲醇，生物除磷后以FeCl₃辅助除磷。")
    add_heading(doc, "3.2 水线流程", 2)
    add_body(doc, "水线采用两条并行线路，最终流程为：粗格栅（地下）→两格进水井/提升泵→细格栅→曝气沉砂池→初沉池→A²/O→二沉配水井→二沉池→接触池（次氯酸钠）→长喉计量渠→汇合井→东厂界控制井。")
    add_figure(doc, "3-1", "图3-1 水线工艺流程及回流、排泥接口")
    add_heading(doc, "3.3 主要设计参数", 2)
    add_table(doc, "表3-2 主要设计参数及属性", ["参数", "采用值", "依据或属性"], [
        ["Kz", "1.50", "GB 50014—2021第4.1.4条及表4.1.15"],
        ["最低设计温度", "15 ℃", "任务书冬季平均温度"],
        ["MLSS / MLVSS", "3.50 / 2.45 kg/m³", "HJ 576—2010；MLVSS/MLSS=0.70"],
        ["SRT", "20 d", "HJ 576—2010和GB 50014—2021范围"],
        ["RAS / 内回流", "0.75Q / 2.50Q", "运行按泥位和NOx反馈调节"],
        ["A²/O HRT", "1.5+4.8+7.7=14.0 h", "缺氧段按15 ℃反硝化能力确定"],
        ["设计出水", "BOD₅/SS/TN/NH₄-N/TP=20/20/20/8/1 mg/L", "课程任务书"],
    ])


def add_chapter_4(doc: Document) -> None:
    add_heading(doc, "4 格栅", 1, True)
    add_heading(doc, "4.1 布置与流量", 2)
    add_body(doc, "粗、细格栅均设两条独立水线，每线1台机械格栅，单台设计流量0.9375 m³/s。单线检修时限流运行，或启用经水力验算的备用、超越通道；现有单线尺寸不宣称可独立承担全厂最高时流量。")
    add_heading(doc, "4.2 过栅流速、堵塞与水损", 2)
    add_formula(doc, "v=Q/[bnh/√(sinα)]，α=60°，正常水深h=0.80 m")
    add_table(doc, "表4-1 栅隙、净空隙数与渠道尺寸汇总", ["项目", "粗格栅", "细格栅"], [
        ["栅隙/净空隙数", "25 mm / 48", "10 mm / 122"],
        ["渠道净宽", "1.75 m", "2.10 m"],
        ["单台流量", "0.9375 m³/s", "0.9375 m³/s"],
        ["正常过栅流速", "0.909 m/s", "0.894 m/s"],
        ["25%堵塞、v=1.0所需水深", "0.969 m", "0.953 m"],
    ])
    add_table(doc, "表4-2 栅条、槽长与水头损失汇总", ["项目", "粗格栅", "细格栅"], [
        ["栅条厚度", "10 mm", "7 mm"],
        ["栅条组合宽度", "1.670 m", "2.067 m"],
        ["水力总长", "2.14 m", "2.96 m"],
        ["清洁水头损失", "0.019 m", "0.039 m"],
        ["设计水头损失", "0.038 m", "0.078 m"],
    ])
    add_heading(doc, "4.3 栅渣量与设备接口", 2)
    add_body(doc, "25 mm粗格栅湿栅渣按0.05 m³/10³m³取5.4 m³/d；10 mm细格栅在粗格栅后的增量暂按同值作设备能力上界，两级合计10.8 m³/d。按湿密度960 kg/m³折合约10.4 t/d，此值为湿栅渣而非干固体。两级均配密闭输送压榨设备。")


def add_chapter_5(doc: Document) -> None:
    add_heading(doc, "5 进水井与提升泵", 1, True)
    add_heading(doc, "5.1 进水井及泵组配置", 2)
    add_body(doc, "进水井设两格，每格净平面16 m×10 m、调节水深1.20 m，有效容积192 m³；每格设3个泵位，按2用1备运行。全厂6台同型潜污泵，最高时4台工作，两条水线各2台工作。")
    add_formula(doc, "Vmin=Qp×5 min=0.46875×5×60=140.625 m³")
    add_formula(doc, "Vcell=16×10×1.20=192 m³>140.625 m³")
    add_body(doc, "单泵额定流量1687.5 m³/h（0.46875 m³/s）。平均日每线2250 m³/h时，1台连续运行、第2台按液位间歇启停；最高时每线2台全开。不采用变频调速，每小时启动次数控制不大于6次。")
    add_figure(doc, "5-1", "图5-1 两格进水井与六泵位计算图")
    add_heading(doc, "5.2 扬程与电机功率", 2)
    add_body(doc, "最不利工况采用进水井最低水位EL22.60 m和细格栅前EL34.80 m。单泵DN700支管按17 m计算，单线DN1000出水管按11 m平面长度加2 m设备连接计算，λ=0.020。")
    add_formula(doc, "Hreq=(34.80−22.60)+0.371057+0.306825+0.80=13.677882 m")
    add_formula(doc, "Pshaft=ρgQpH/(η×1000)=1000×9.81×0.46875×14/(0.80×1000)=80.47 kW")
    add_body(doc, "所需扬程13.677882 m，采用14 m；轴功率加10%余量为88.52 kW，选90 kW额定轴输出电机。电输入功率还应除以电机效率。采购时应由厂家复核全工况曲线、NPSH、淹没深度和启停频率。")
    add_table(doc, "表5-1 提升泵扬程组成及设备配置", ["项目", "采用值", "说明"], [
        ["单泵设计流量", "1687.5 m³/h", "每线2台工作"],
        ["静扬程", "12.20 m", "34.80−22.60"],
        ["DN700支管损失", "0.371057 m", "L=17 m"],
        ["DN1000出水管损失", "0.306825 m", "L=13 m"],
        ["未细化余量", "0.80 m", "阀件、渐变及设备接口"],
        ["所需/采用扬程", "13.677882 / 14.0 m", "不重复计入后续重力落差"],
        ["泵组与电机", "6台4用2备；90 kW", "两条水线各2用1备"],
    ])


def add_chapter_6(doc: Document) -> None:
    # 第5章末表格会自然推至下一页；此处再强制分页会在 Word 中生成空白页。
    add_heading(doc, "6 曝气沉砂池", 1)
    add_heading(doc, "6.1 池体尺寸与水力校核", 2)
    add_body(doc, "采用2座矩形曝气沉砂池，单池按最高时单线流量0.9375 m³/s校核。净尺寸L×B×h=30.0 m×3.80 m×2.55 m，单池有效容积290.70 m³。")
    add_formula(doc, "v=0.9375/(3.80×2.55)=0.0967 m/s")
    add_formula(doc, "tmax=290.70/0.9375/60=5.168 min；tavg=7.752 min")
    add_table(doc, "表6-1 曝气沉砂池校核", ["项目", "计算值", "采用或要求"], [
        ["宽深比B/h", "1.49", "1.0～1.5"],
        ["长宽比L/B", "7.89", "满足课程布置"],
        ["最高时水平流速", "0.0967 m/s", "不大于0.10 m/s"],
        ["最高时停留时间", "5.168 min", "大于5 min"],
        ["平均日停留时间", "7.752 min", "运行参考"],
    ])
    add_figure(doc, "6-1", "图6-1 曝气沉砂池计算图")
    add_heading(doc, "6.2 供气、排砂与高程接口", 2)
    add_body(doc, "线供气量按GB 50014—2021第7.4节规定范围5～12 L/(m·s)选取，单池沿30 m池长布气；实际风量由曝气强度与设备样本复核。池底采用平底，主体水位EL34.25 m、平底EL31.70 m，不设池内局部砂斗。吸砂机沿池长往复排砂，砂水混合物送池外砂水分离器。")
    add_body(doc, "按0.03 m³砂/1000 m³污水的课程取值，平均每池砂量1.62 m³/d，两日为3.24 m³。池外储砂斗几何容积3.26 m³、有效工作容积3.20 m³，对应约1.98 d砂量；设计按每日排砂运行。")


def add_chapter_7(doc: Document) -> None:
    add_heading(doc, "7 初次沉淀池", 1, True)
    add_heading(doc, "7.1 池数、面积及水力校核", 2)
    add_body(doc, "初沉采用4座平流式沉淀池，两条水线各2座。按最高时流量6750 m³/h和表面负荷2.007 m³/(m²·h)反算总净面积3364.0 m²、单池841.0 m²；取单池净宽14.5 m，净长58.0 m，有效水深3.5 m。")
    add_formula(doc, "qmax=6750/(4×58×14.5)=2.006 m³/(m²·h)")
    add_formula(doc, "tmax=58×14.5×3.5/(6750/4)=1.744 h")
    add_formula(doc, "vH=0.46875/(14.5×3.5)=0.00924 m/s")
    add_table(doc, "表7-1 初沉池水力校核", ["工况或指标", "计算值", "结论"], [
        ["平均日表面负荷", "1.338 m³/(m²·h)", "运行参考"],
        ["最高时表面负荷", "2.006 m³/(m²·h)", "设计主算"],
        ["平均/最高时停留时间", "2.616 / 1.744 h", "峰时满足"],
        ["最高时水平流速", "0.00924 m/s", "位于常用范围"],
        ["长宽比/长深比", "4.0 / 16.57", "满足常用要求"],
    ])
    add_figure(doc, "7-1", "图7-1 平流式初次沉淀池平剖计算图")
    add_heading(doc, "7.2 出水堰、刮泥与泥斗", 2)
    add_body(doc, "每池设12条指形出水槽，每条有效堰长14 m，总堰长168 m。最高时单池流量0.46875 m³/s，堰负荷为2.790 L/(m·s)，低于GB 50014—2021第7.5.8条2.9 L/(m·s)上限。")
    add_body(doc, "池底按0.01纵坡坡向进水端的贯通条形泥斗。泥斗上口沿池宽14.5 m贯通、纵向宽3.0 m，下口同样沿池宽贯通、纵向宽0.8 m，深2.0 m；梯形棱柱容积为14.5×(3.0+0.8)/2×2.0=55.10 m³，大于单池4 h贮泥需求15.625 m³。主体水位EL33.05 m，名义沉淀底EL29.55 m，泥斗端坡底EL28.97 m，斗底EL26.97 m。")
    add_heading(doc, "7.3 初沉出水负荷", 2)
    add_table(doc, "表7-2 初沉去除及出水负荷", ["指标", "去除率", "出水负荷 kg/d", "出水浓度 mg/L"], [
        ["BOD₅", "25.0%", "13500", "125.00"],
        ["SS", "50.0%", "15000", "138.89"],
        ["TN", "0%", "4800", "44.44"],
        ["TP", "7.5%", "555", "5.14"],
    ])
    add_body(doc, "TN按初沉基本不去除计入，使后续脱氮容量采用最不利上界；初沉污泥按4%含固率另计。")


def add_chapter_8(doc: Document) -> None:
    add_heading(doc, "8 A²/O生物反应池与回流", 1, True)
    add_heading(doc, "8.1 容积与分区", 2)
    add_body(doc, "设置4个工艺系列，按两条水线成对布置并共享备用接口。每系列净水面70 m×45 m、净水深5 m、有效容积15750 m³；总有效容积63000 m³。每系列6条7.5 m净宽廊道，等效流长420 m。")
    add_table(doc, "表8-1 A²/O分区与水力停留时间", ["分区", "HRT h", "全厂容积 m³", "单系列容积 m³", "等效长度 m/系列"], [
        ["厌氧", "1.50", "6750", "1687.5", "45"],
        ["缺氧", "4.80", "21600", "5400", "144"],
        ["好氧", "7.70", "34650", "8662.5", "231"],
        ["合计", "14.00", "63000", "15750", "420"],
    ])
    add_figure(doc, "8-1", "图8-1 A²/O单系列分区与回流泵接口图")
    add_heading(doc, "8.2 污泥浓度、负荷和SRT", 2)
    add_body(doc, "基准MLSS取3.50 kg/m³，MLVSS/MLSS取0.70，故MLVSS=2.45 kg/m³；系统SRT取20 d。初沉后BOD₅负荷13500 kg/d，达到20 mg/L目标需去除11340 kg/d。")
    add_formula(doc, "LMLSS=13500/(63000×3.50)=0.0612 kgBOD₅/(kgMLSS·d)")
    add_formula(doc, "LMLVSS=13500/(63000×2.45)=0.0875 kgBOD₅/(kgMLVSS·d)")
    add_body(doc, "采用值与GB 50014—2021表7.6.19给出的A²/O总HRT、分区HRT、MLSS、SRT及内回流范围相符。厌氧区DO控制小于0.2 mg/L，缺氧区0.2～0.5 mg/L，好氧区不低于2 mg/L。")
    add_heading(doc, "8.3 回流系统", 2)
    add_table(doc, "表8-2 RAS与内回流设计流量", ["系统", "全厂平均 m³/h", "全厂最高时 m³/h", "最高时单系列 m³/h"], [
        ["回流污泥 R=0.75Q", "3375", "5062.5", "1265.625"],
        ["混合液内回流 Ri=2.50Q", "11250", "16875", "4218.75"],
    ])
    add_table(doc, "表8-3 回流泵配置", ["设备", "单台设计点", "数量与用备", "管径", "扬程/电机"], [
        ["回流污泥泵", "1266 m³/h", "6台，4用2备", "单池DN500、系列DN700", "3.0 m / 18.5 kW"],
        ["混合液内回流泵", "4219 m³/h", "6台，4用2备", "系列DN1200", "1.2 m / 22 kW"],
    ])
    add_body(doc, "每条水线服务2个A²/O系列，并设2台工作泵和1台共享备用泵；备用泵通过共用取水廊道、切换阀组和止回设施择一接入同线两个系列。RAS按二沉泥位和MLSS调节，内回流按缺氧区NOx反馈调节。")
    add_heading(doc, "8.4 低温硝化边界", 2)
    add_body(doc, "任务书未给进水TKN及NH₄-N连续数据，因此不建立来源不足的AOB动力学结论。以TN负荷减出水NH₄-N限值所得3936 kgN/d作为硝化设备容量上界；系统SRT=20 d位于HJ 576—2010和GB 50014—2021推荐范围。实际工程须以最低连续水温、TKN、NH₄-N、pH、碱度和抑制物重新复核。")


def add_chapter_9(doc: Document) -> None:
    add_heading(doc, "9 脱氮碳源、除磷与碱度", 1, True)
    add_heading(doc, "9.1 反硝化容积", 2)
    add_body(doc, "按HJ 576—2010式(8)～(10)计算微生物量和缺氧池容积。由于缺少TKN，暂以TN=44.44 mg/L代替Nk，并明确为课程容量上界。")
    add_formula(doc, "ΔXv=Q(S0−Se)yYt/1000=108000×(125−20)×0.70×0.40/1000=3175.2 kg/d")
    add_formula(doc, "Kde,15=0.045×1.08^(15−20)=0.03063 kgN/(kgMLSS·d)")
    add_formula(doc, "Vn=[0.001Q(Nk−Nte)−0.12ΔXv]/(Kde,15X)=21074 m³")
    add_body(doc, "采用缺氧区21600 m³，15 ℃反硝化能力2315.34 kgN/d，较需反硝化量2258.98 kgN/d留有2.50%余量。")
    add_heading(doc, "9.2 外加碳源", 2)
    add_body(doc, "以原水可利用反硝化碳贡献为零作为设备容量边界。甲醇反硝化计量关系采用5 mol CH₃OH/6 mol NO₃-N，即纯甲醇1.9048 kg/kgN；商品液质量分数80%，并加20%课程容量余量。")
    add_formula(doc, "M80=2258.98×1.9048/0.80×1.20=6454 kg/d，采用6500 kg/d")
    add_body(doc, "峰时能力为403.4 kg/h，设置2台100%容量计量泵，单台不小于450 kg/h。6500 kg/d为设备容量，不是固定运行量；运行按NOx、TN和甲醇残余反馈调节。")
    add_heading(doc, "9.3 化学辅助除磷", 2)
    add_body(doc, "初沉后TP为555 kgP/d。按生物TP去除率60%下限，A²/O后剩余222 kgP/d；达到108 kgP/d目标需化学去除114 kgP/d。初算Fe/P=1.50 mol/mol。")
    add_formula(doc, "纯FeCl₃=114×1.50×162.2/31=895.5 kg/d")
    add_body(doc, "40% FeCl₃溶液密度按1.40 t/m³，计算量1.599 m³/d，加20%容量后采用2.0 m³/d；设置2台100%计量泵，单台不小于0.15 m³/h。化学干固体约751.747 kg/d。实际药剂投加量须用烧杯试验或同类工程数据修正。")
    add_figure(doc, "9-1", "图9-1 投药与固体闭合边界图")
    add_heading(doc, "9.4 碱度与总固体", 2)
    add_body(doc, "硝化耗碱按7.14 kgCaCO₃/kgN，反硝化回补按3.57 kgCaCO₃/kgN，铁盐水解按3当量/mol Fe初算，并保留出水残余碱度70 mg/L。以TN代NH₄-N的上界所需进水碱度为263.21 mg/L。")
    add_body(doc, "若实测进水碱度为100 mg/L，99% NaHCO₃需求29.91 t/d，加20%能力后为35.9 t/d。该值仅为设备容量边界，获得实测碱度后应按缺口调整。")
    add_table(doc, "表9-1 固体库存与排泥", ["项目", "计算", "结果"], [
        ["基准固体库存", "63000×3.50", "220500 kg"],
        ["20 d基准排固", "220500/20", "11025 kg/d"],
        ["化学固体库存", "751.747×20", "15034.947 kg"],
        ["总MLSS", "3.50+15034.947/63000", "3.738650 kg/m³"],
        ["总排固", "63000×3.738650/20", "11776.747 kg/d"],
    ])


def add_chapter_10(doc: Document) -> None:
    add_heading(doc, "10 需氧量、曝气器与鼓风机", 1, True)
    add_heading(doc, "10.1 需氧量", 2)
    add_body(doc, "按HJ 576—2010式(15)，外加80%甲醇6500 kg/d中纯甲醇为5200 kg/d，按1.50 kgCOD/kg甲醇折合7800 kgCOD/d。反硝化氧当量只扣除一次。")
    add_formula(doc, "O₂=1.47×11340+7800−1.42×3175.2+4.57×(3936−381.024)−0.62×4.57×2258.976")
    add_table(doc, "表10-1 需氧量分项", ["分项", "kgO₂/d"], [
        ["BOD₅碳氧化", "16669.800"],
        ["外加甲醇COD容量", "7800.000"],
        ["微生物量氧当量抵扣", "−4508.784"],
        ["硝化", "16246.240"],
        ["反硝化氧当量抵扣", "−6400.583"],
        ["合计AOR", "29806.674"],
    ])
    add_formula(doc, "AORdesign=(29806.674/24)×1.50×1.10=2049.209 kgO₂/h")
    add_heading(doc, "10.2 标准供气量和曝气器", 2)
    add_body(doc, "采用α=0.80、β=0.95、25 ℃、好氧DO=2 mg/L、曝气器淹没4.75 m，现场/标准转移系数为0.671220，故SOR=3052.962 kgO₂/h。229 mm盘式微孔曝气器样本在该淹没深度下的清水SOTE包络约28.5%～38%，设计取EA=25%考虑污水、水龄和布置影响。")
    add_formula(doc, "Gs,20=3052.962/(0.28×0.25)=43613.746 m³/h（20 ℃）")
    add_formula(doc, "GN=43613.746×273.15/293.15=40638.222 Nm³/h，采用48000 Nm³/h")
    add_body(doc, "每只曝气器按5.0 Nm³/h配置，共9600只、每系列2400只。好氧区总底面积6930 m²，布置密度1.385只/m²。采用风量折回20 ℃后可提供2420.431 kgO₂/h，较设计AOR留有18.12%余量。")
    add_figure(doc, "10-1", "图10-1 曝气与鼓风系统计算图")
    add_heading(doc, "10.3 鼓风机和风压", 2)
    add_table(doc, "表10-2 鼓风机压力组成", ["压力项", "kPa"], [
        ["4.75 m淹没静压", "46.60"],
        ["曝气器", "5.00"],
        ["管网与阀件", "6.00"],
        ["余量", "3.00"],
        ["合计", "60.60"],
        ["采用压力", "65.00(g)"],
    ])
    add_body(doc, "设6台高速离心鼓风机、4用2备，单台12000 Nm³/h、65 kPa(g)，4台工作总风量48000 Nm³/h。空气系统采用DN1200环状母管和4根DN700系列支管；环路一段隔离时最不利外管损失2.523 kPa，小于6 kPa管阀预算，剩余3.477 kPa用于池内配气。最终工作点、效率、轴功率、喘振及噪声由厂家曲线确认。")


def add_chapter_11(doc: Document) -> None:
    add_heading(doc, "11 二次沉淀池与固体平衡", 1, True)
    add_heading(doc, "11.1 几何与水力负荷", 2)
    add_body(doc, "采用8座辐流二沉池，两条水线各4座。每池净径37 m、中心井D6 m，扣除中心井后总净面积为8375.486 m²。沉淀区有效水深3.1 m，缓冲层0.50 m、水平基准贮泥层0.40 m，周边总深4.00 m；环形底坡0.05形成0.775 m坡降，中心集泥坑深0.50 m，中心总深5.275 m。")
    add_formula(doc, "Anet=8×π(37²−6²)/4=8375.486 m²")
    add_table(doc, "表11-1 二沉池水力与固体负荷校核", ["校核项", "计算值", "控制值或结论"], [
        ["最高时表面负荷", "0.806 m³/(m²·h)", "满足课程采用范围"],
        ["最高时外部HRT", "3.847 h", "满足1.5～4 h"],
        ["D/h", "11.94", "满足6～12"],
        ["最高时固体负荷", "126.683 kg/(m²·d)", "小于150"],
        ["最高时堰负荷", "1.130 L/(m·s)", "小于1.7"],
    ])
    add_figure(doc, "11-1", "图11-1 二次沉淀池平剖计算图")
    add_heading(doc, "11.2 固体负荷与排泥", 2)
    add_body(doc, "投药点位于A²/O出水后、二沉池前。进入二沉的总MLSS为3.738650 kg/m³；随流量同比调节的当次新生成化学固体在进口增加一次。")
    add_formula(doc, "SLRmax=[162000×1.75×3.738650+1.50×751.747]/8375.486=126.683 kg/(m²·d)")
    add_body(doc, "每池有效贮泥由0.40 m水平层418.77 m³、环形坡底楔体308.20 m³及D6 m、深0.50 m集泥坑14.14 m³组成，合计741.11 m³。二沉连续刮吸泥，低流量时按泥位减少运行池数。")
    add_heading(doc, "11.3 进出水和堰", 2)
    add_body(doc, "每池进水混合液支管DN900，最高时0.410156 m³/s；出水支管DN700，0.234375 m³/s。中心线D33 m双侧环形堰长207.35 m，最高时堰负荷1.130 L/(m·s)，满足控制要求。主体水位EL31.15 m，周边工艺底EL27.15 m，坡底内缘EL26.375 m，坑底EL25.875 m。")


def add_chapter_12(doc: Document) -> None:
    add_heading(doc, "12 连接管渠与配水设施", 1, True)
    add_heading(doc, "12.1 计算方法与流量口径", 2)
    add_body(doc, "压力管沿程损失采用Darcy-Weisbach式，主水线λ=0.025并以0.030作敏感性复核；局部损失按构造取Σζ。管线长度取总平面坐标折线的中心线长度，另计竖管、设备连接和构筑物内部预算。")
    add_formula(doc, "A=πD²/4；v=Q/A；hf=λ(L/D)v²/(2g)；hl=Σζv²/(2g)")
    add_body(doc, "明渠采用Manning式hf=L(nv/R^(2/3))²，n=0.013。构筑物与提升后连接管渠均按最高时流量6750 m³/h校核；日污染负荷不乘Kz。")
    add_figure(doc, "12-1", "图12-1 主要连接管渠与配水接口图")
    add_heading(doc, "12.2 主水线管径与中心线长度", 2)
    add_table(doc, "表12-1 主水线管径与流量", ["管段", "规格", "最高时单元流量", "代表性中心线长度"], [
        ["泵站至细格栅", "2×DN1000", "0.9375 m³/s/线", "11 m；设备连接另计2 m"],
        ["沉砂至初沉配水", "DN1200", "0.9375 m³/s/线", "A/B线100/70 m"],
        ["初沉配水至单池", "DN900", "0.46875 m³/s/池", "井后4 m+独立支管64 m"],
        ["初沉至A²/O", "DN900", "0.46875 m³/s/系列", "28 m"],
        ["A²/O至二沉配水", "DN1200", "0.820313 m³/s/系列", "58 m"],
        ["二沉配水至单池", "DN900", "0.410156 m³/s/池", "58～98 m"],
        ["二沉出水支管", "DN700", "0.234375 m³/s/池", "3或5 m"],
        ["二沉出水汇流", "DN1200", "0.9375 m³/s/线", "A线263 m、B线166 m"],
        ["汇合井至东厂界", "DN1500", "1.875 m³/s", "48 m"],
    ])
    add_body(doc, "不重复物理段统计的水线中心线总长为2129 m，最不利外部路径为C1—B1—T1，中心线长826 m。该路径由不同流量、管径和构造段组成，不作为单一管径一次计算。")
    add_heading(doc, "12.3 配水井及均匀分配", 2)
    add_body(doc, "初沉每条水线设1座配水井；二沉采用2座配水井，每井接收2个A²/O系列并服务4座二沉池。内回流在A²/O内部循环，不计入二沉外部配水流量；二沉混合液为外部水量与RAS之和。")
    add_table(doc, "表12-2 二沉配水流量与管径", ["位置", "数量", "最高时单元流量", "管径", "流速"], [
        ["A²/O系列至配水井", "4", "0.820313 m³/s", "DN1200", "0.725 m/s"],
        ["每井短公共段", "2", "1.640625 m³/s", "DN1500", "0.928 m/s"],
        ["配水井至二沉池", "8", "0.410156 m³/s", "DN900", "0.645 m/s"],
        ["二沉出水支管", "8", "0.234375 m³/s", "DN700", "0.609 m/s"],
        ["每线出水汇流", "2", "0.9375 m³/s", "DN1200", "0.829 m/s"],
    ])
    add_formula(doc, "h=[Q/(1.84b)]^(2/3)=[0.410156/(1.84×1.20)]^(2/3)=0.3256 m")
    add_body(doc, "每井设置4个1.20 m等高矩形薄壁堰和4条独立DN900支路，堰上水头采用0.35 m预算。两井设联络与隔离；单井检修时限流运行。")
    add_heading(doc, "12.4 进厂管边界", 2)
    add_body(doc, "D1200进厂管由任务书确定，管底EL23.30 m。按h/D=0.75、Manning n=0.013计算，过水面积0.910 m²、湿周2.513 m、水力半径0.362 m，最高时流速2.06 m/s、坡度2.78‰。施工深化应核对上游水位、管材允许流速和壅水风险。")


def add_chapter_13(doc: Document) -> None:
    add_heading(doc, "13 总平面布置", 1, True)
    add_heading(doc, "13.1 坐标、功能分区及流程", 2)
    add_body(doc, "厂区按400 m×300 m矩形表示12 hm²，西南角为坐标原点，x向东、y向北。北界进水点为(360,300)，东厂界出水点为(400,29)。流程从东北入厂、向西分流、向南处理并折返向东出厂。")
    add_table(doc, "表13-1 主要构筑物总平面接口", ["单元", "数量", "中心坐标或排列", "工艺净尺寸"], [
        ["粗格栅", "1座2格", "(360,282)", "2格"],
        ["进水井/泵站", "1座2格", "(360,262)", "2×(16×10×1.2 m)"],
        ["曝气沉砂池", "2座", "(207,274)、(257,274)", "30×3.8×2.55 m/座"],
        ["平流初沉池", "4座", "x=50/150/250/350，y=226", "58×14.5×3.5 m/座"],
        ["A²/O", "4系列", "x=50/150/250/350，y=163", "70×45×5 m/系列"],
        ["二沉配水井", "2座", "(100,124)、(300,124)", "每井服务4池"],
        ["辐流二沉池", "8座", "x=34～366，y=80", "净D37、中心井D6"],
        ["折流接触池", "1座2格", "(299,28.5)", "2×(45×8×4.8 m)"],
        ["计量渠", "2条", "(338,29)", "喉宽1.50 m"],
    ])
    add_heading(doc, "13.2 道路、间距和面积", 2)
    add_body(doc, "设置6 m厂界环路和三条6 m横向联系路，南侧设4 m次通道；南侧主门服务管理和日常交通，北侧服务门用于设备检修和药剂运输。管理区靠南侧主门独立布置，与前处理、加药、污泥和鼓风机房之间设置绿化隔离。任务书只给夏季东南风，没有全年风频和静风率，因此不作已验证上风向结论。")
    add_table(doc, "表13-2 用地面积平衡", ["分项", "面积 m²"], [
        ["水处理构筑物包络", "37990"],
        ["辅助及管理设施预留", "5008"],
        ["连续发展用地", "2800"],
        ["道路并集", "16240"],
        ["未分配用地", "57962"],
        ["合计", "120000"],
    ])
    add_body(doc, "未分配用地中列36000 m²绿化目标，其余用于管廊、除臭、消防、结构修正和辅助设施。构筑物总图包络用于课程排布，不等同于结构外包尺寸。")
    add_figure(doc, "13-1", "图13-1 厂区总平面与管线图")
    add_heading(doc, "13.3 检修与扩建", 2)
    add_body(doc, "前处理两线、进水井两格、生物池四系列和二沉八池均可分单元隔离。单线或单格检修时按剩余能力限流，不宣称未经核算的单线承担全厂峰值。北西侧预留连续发展用地；事故超越只保留受控接口，任何排放须依法审批。")


HYDRAULIC_ROWS = [
    ("细格栅", "0.15000", "0.15", "34.80", "34.65"),
    ("细栅后至沉砂入口", "0.06331", "0.10", "34.65", "34.55"),
    ("沉砂池内", "0.25100", "0.30", "34.55", "34.25"),
    ("沉砂至初沉配水", "0.37824", "0.45", "34.25", "33.80"),
    ("初沉自由配水", "0.35316", "0.40", "33.80", "33.40"),
    ("初沉支管及进水", "0.29022", "0.35", "33.40", "33.05"),
    ("初沉出水堰槽", "0.20000", "0.20", "33.05", "32.85"),
    ("初沉集水槽至A²/O入口", "0.20569", "0.25", "32.85", "32.60"),
    ("A²/O入口至主体", "0.30000", "0.30", "32.60", "32.30"),
    ("A²/O至二沉配水", "0.26612", "0.35", "32.30", "31.95"),
    ("二沉自由配水", "0.42556", "0.45", "31.95", "31.50"),
    ("二沉支管及中心进水", "0.26529", "0.35", "31.50", "31.15"),
    ("二沉出水堰槽", "0.20000", "0.20", "31.15", "30.95"),
    ("二沉出水至接触池", "0.60571", "0.70", "30.95", "30.25"),
    ("接触池廊道及整流", "0.10256", "0.15", "30.25", "30.10"),
    ("接触池后至计量上游", "0.01118", "0.05", "30.10", "30.05"),
    ("长喉计量与尾端消能", "0.75000", "0.75", "30.05", "29.30"),
    ("计量后至汇合井", "0.26748", "0.35", "29.30", "28.95"),
    ("汇合井至东厂界", "0.14345", "0.20", "28.95", "28.75"),
]


def add_chapter_14(doc: Document) -> None:
    add_heading(doc, "14 水力高程与出厂方式", 1, True)
    add_heading(doc, "14.1 高程计算原则", 2)
    add_body(doc, "以A²/O主体内底EL27.30 m和水位EL32.30 m为竖向锚点，按最不利路径分别向上游和下游推算。连接管采用实际中心线长度、竖管预算和局部构造损失；配水堰、池内整流和计量消能分别计入，避免把经验落差替代计算。")
    add_figure(doc, "14-1", "图14-1 全厂水线高程控制图")
    add_heading(doc, "14.2 逐段水头损失与采用落差", 2)
    add_table(doc, "表14-1 重力水线逐段高程", ["段落", "基准需求 m", "采用落差 m", "上游EL m", "下游EL m"], HYDRAULIC_ROWS)
    add_body(doc, "细格栅前至计量后落差为34.80−29.30=5.50 m；延伸到东厂界控制井后为34.80−28.75=6.05 m。最不利C1—B1—T1路径基准需求5.228250 m，采用落差中有0.821750 m用于调节和余量。各段余量只在本段通过阀门、堰和出口控制消耗。")
    add_heading(doc, "14.3 主要节点高程", 2)
    add_table(doc, "表14-2 主要节点水位和工艺内底", ["节点", "水位EL m", "工艺内底EL m", "说明"], [
        ["进厂D1200北界", "峰24.200", "管底23.300", "地面27.300"],
        ["进水井", "22.600～23.800", "20.400", "局部泵坑19.800"],
        ["细格栅前/后", "34.800/34.650", "后渠33.650", "平台35.300"],
        ["沉砂主体", "34.250", "31.700", "平底，无池内砂斗"],
        ["初沉主体", "33.050", "名义29.550", "坡底28.970；斗底26.970"],
        ["A²/O入口/主体", "32.600/32.300", "27.300", "主体净深5.00 m"],
        ["二沉配水上/下", "31.950/31.500", "29.950/29.500", "2井×4池"],
        ["二沉主体", "31.150", "周边27.150", "坑底25.875"],
        ["接触池入口/出口", "30.250/30.100", "25.300", "出口峰时水深4.80 m"],
        ["计量上游/后井", "30.050/29.300", "29.050/28.300", "喉底29.527691"],
        ["汇合井/东厂界", "28.950/28.750", "24.500", "DN1500至厂界"],
    ])
    add_heading(doc, "14.4 泵扬程和出厂边界", 2)
    add_body(doc, "提升泵所需扬程为13.677882 m，采用14 m；6.05 m重力落差已通过细格栅前标高体现，不再重复叠加。RAS最不利所需2.521818 m，采用3.0 m；内回流所需0.914649 m，采用1.2 m；风机采用65 kPa(g)。")
    add_body(doc, "厂内DN1500至东厂界48 m已计入水线。厂外尚无实测，按100 m、DN1500、λ=0.025、Σζ=3.0进行条件算例，含0.10 m余量时允许尾水上限EL28.382 m。任务书给出的洪水EL25.00 m下具重力出厂余量；组号公式18+0.5M仍须由提交人用实际组号复核。")


def add_chapter_15(doc: Document) -> None:
    add_heading(doc, "15 接触消毒与计量", 1, True)
    add_heading(doc, "15.1 接触池容积与几何", 2)
    add_body(doc, "采用2格折流接触池，每格净水尺寸45 m×8 m×4.8 m、有效容积1728 m³，总有效容积3456 m³。每格设4条2.0 m净宽廊道，等效流程约180 m。")
    add_formula(doc, "Vreq=6750×30/60=3375 m³；tmax=3456/6750×60=30.72 min")
    add_body(doc, "两格正常同时运行；单格检修时另一格不具备全峰值30 min容积，应限流或配置经验证的替代消毒措施。若3道隔墙厚度暂按0.20 m，结构池内总宽至少8.60 m，不能把8.0 m净水宽理解为含墙总宽。")
    add_figure(doc, "15-1", "图15-1 接触消毒与计量终端图")
    add_heading(doc, "15.2 次氯酸钠设备能力", 2)
    add_body(doc, "二沉澄清出水在接触池前投加次氯酸钠。无消毒试验资料时，按GB 50014—2021第7.13节采用5 mg/L有效氯作为名义投加量、15 mg/L作为设备能力上界；运行剂量由流量前馈、余氯反馈和微生物检测确定。")
    add_table(doc, "表15-1 次氯酸钠设计能力", ["工况", "有效氯", "10% NaOCl商品液"], [
        ["平均日、5 mg/L", "540 kg/d", "4.724 m³/d"],
        ["平均日、15 mg/L", "1620 kg/d", "14.173 m³/d"],
        ["最高时、5 mg/L", "33.75 kg/h", "0.295 m³/h"],
        ["最高时、15 mg/L", "101.25 kg/h", "0.886 m³/h"],
    ])
    add_body(doc, "设置2台100%隔膜计量泵、1用1备，单台1.0 m³/h；储罐按平均日15 mg/L、7 d上限确定工作容积99.21 m³，采用2座约50 m³储罐，低温、避光储存。商品液纯度和有效氯含量以供货证书为准。")
    add_heading(doc, "15.3 长喉计量渠", 2)
    add_body(doc, "接触池后设2条长喉计量渠，每条按0.9375 m³/s设计，喉宽1.50 m、课程选喉长2.5 m。")
    add_formula(doc, "yc=[Q²/(gb²)]^(1/3)=0.341479 m")
    add_body(doc, "计量上游水位EL30.05 m，喉底EL29.527691 m，临界水面EL29.869170 m；计量后水位EL29.30 m，具自由跌水边界。临界能量公式用于课程水力初算，最终流量关系须按选定槽型及厂家标定曲线校核。")


def add_chapter_16(doc: Document) -> None:
    add_heading(doc, "16 主要构筑物和设备汇总", 1, True)
    add_heading(doc, "16.1 主要构筑物", 2)
    add_table(doc, "表16-1 主要构筑物汇总", ["单元", "数量", "采用规格", "关键校核"], [
        ["粗格栅", "2线", "b=25 mm，B=1.75 m", "v=0.909 m/s"],
        ["进水井", "2格", "每格16×10×1.2 m，192 m³", "大于单泵5 min水量140.625 m³"],
        ["细格栅", "2线", "b=10 mm，B=2.10 m", "v=0.894 m/s"],
        ["曝气沉砂池", "2座", "30×3.8×2.55 m", "v=0.0967 m/s；t=5.168 min"],
        ["平流初沉池", "4座", "58×14.5×3.5 m", "qmax=2.006；t=1.744 h"],
        ["A²/O", "4系列", "70×45×5 m；总63000 m³", "HRT=14 h；SRT=20 d"],
        ["二沉配水井", "2座", "每井4堰、4条DN900支路", "每井服务4座二沉池"],
        ["辐流二沉池", "8座", "净D37、中心井D6", "SLRmax=126.683"],
        ["折流接触池", "2格", "45×8×4.8 m/格", "tmax=30.72 min"],
        ["长喉计量渠", "2条", "b=1.50 m，喉长2.5 m", "yc=0.341479 m"],
    ])
    add_heading(doc, "16.2 主要设备", 2)
    add_table(doc, "表16-2 主要设备设计工况", ["设备", "数量与备用", "单台或系统设计工况", "控制要求"], [
        ["机械粗格栅", "2台", "栅隙25 mm，Q≥0.9375 m³/s", "液位差控制清渣"],
        ["机械细格栅", "2台", "栅隙10 mm，Q≥0.9375 m³/s", "配密闭压榨输送"],
        ["进水提升泵", "6台，4用2备", "1687.5 m³/h，H=14 m，90 kW", "定速台数/液位调节"],
        ["沉砂池吸砂机", "2套", "沿30 m池长往复", "池外砂水分离"],
        ["初沉刮泥机", "4套", "58 m平流池", "向进水端条形斗刮泥"],
        ["RAS泵", "6台，4用2备", "1266 m³/h，H=3.0 m，18.5 kW", "泥位/MLSS反馈"],
        ["内回流泵", "6台，4用2备", "4219 m³/h，H=1.2 m，22 kW", "NOx反馈"],
        ["微孔曝气器", "9600只", "229 mm盘式，5.0 Nm³/h·只", "每系列2400只"],
        ["高速离心鼓风机", "6台，4用2备", "12000 Nm³/h，65 kPa(g)", "环状母管切换"],
        ["二沉刮吸泥机", "8套", "D37辐流池", "连续排泥"],
        ["甲醇计量泵", "2台100%", "≥450 kg/h·台", "按NOx/TN反馈"],
        ["FeCl₃计量泵", "2台100%", "≥0.15 m³/h·台", "烧杯试验修正"],
        ["NaOCl计量泵", "2台，1用1备", "1.0 m³/h·台", "流量、余氯反馈"],
    ])
    add_heading(doc, "16.3 提交前核验结论", 2)
    add_body(doc, "设计流量、泵组、构筑物数量与尺寸、A²/O分区、回流、供氧、二沉固体负荷、配水、消毒和水力高程已按同一组设计数据闭合。正文、表格、设备汇总和12幅图采用唯一值，文档审阅痕迹已经清除。")
    add_body(doc, "课程设计计算不存在未闭合数值。封面班级、组号和成员由提交人填写；组号还应代入任务书洪水式18+0.5M复核。实际工程仍须补实测水质、最低连续水温、碱度、风玫瑰、地形地勘、厂外排口、消毒与烧杯试验，以及泵和风机厂家曲线。")


def add_references(doc: Document) -> None:
    add_heading(doc, "参考文献", 1, True)
    refs = [
        "[1] 中国农业大学环境工程课程组. 本科生课程设计任务书（6）及污水处理厂课程设计指导书, 2026.",
        "[2] 住房和城乡建设部. GB 55027—2022 城乡排水工程项目规范. 北京: 中国建筑工业出版社, 2022.",
        "[3] 住房和城乡建设部. GB 50014—2021 室外排水设计标准. 北京: 中国计划出版社, 2021.",
        "[4] 生态环境部. GB 18918—2002 城镇污水处理厂污染物排放标准及2025修改单, 2025.",
        "[5] 江苏省市场监督管理局, 江苏省生态环境厅. DB32/4440—2022 城镇污水处理厂污染物排放标准, 2022.",
        "[6] 环境保护部. HJ 576—2010 厌氧—缺氧—好氧活性污泥法污水处理工程技术规范, 2010.",
        "[7] 给水排水设计手册编写组. 给水排水设计手册 第1册 常用资料. 北京: 中国建筑工业出版社.",
        "[8] 给水排水设计手册编写组. 给水排水设计手册 第5册 城镇排水. 北京: 中国建筑工业出版社.",
        "[9] 中国农业大学资源与环境学院. 环境工程设计教案（污水处理厂课程设计）, 2026.",
        "[10] Xylem. Sanitaire Silver Series II Diffuser Datasheet, Standard 229 mm.",
        "[11] AERZEN. AERZEN Turbo Generation 5plus product information.",
        "[12] U.S. Bureau of Reclamation. Water Measurement Manual, long-throated flumes.",
    ]
    for ref in refs:
        add_body(doc, ref, indent=False)


def strip_comments_and_personal_metadata(path: Path) -> None:
    tmp = path.with_suffix(".tmp.docx")
    remove_parts = {
        "word/comments.xml",
        "word/commentsExtended.xml",
        "word/people.xml",
        "word/person.xml",
    }
    with zipfile.ZipFile(path, "r") as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            if item.filename in remove_parts:
                continue
            data = zin.read(item.filename)
            if item.filename == "docProps/core.xml":
                text = data.decode("utf-8")
                text = re.sub(r"<dc:creator>.*?</dc:creator>", "<dc:creator></dc:creator>", text)
                text = re.sub(r"<cp:lastModifiedBy>.*?</cp:lastModifiedBy>", "<cp:lastModifiedBy></cp:lastModifiedBy>", text)
                data = text.encode("utf-8")
            if item.filename.endswith(".rels"):
                text = data.decode("utf-8")
                text = re.sub(r'<Relationship[^>]+Type="[^"]+/comments(?:Extended)?"[^>]*/>', "", text)
                data = text.encode("utf-8")
            if item.filename == "[Content_Types].xml":
                text = data.decode("utf-8")
                text = re.sub(r'<Override[^>]+PartName="/word/comments(?:Extended)?\.xml"[^>]*/>', "", text)
                text = re.sub(r'<Override[^>]+PartName="/word/(?:people|person)\.xml"[^>]*/>', "", text)
                data = text.encode("utf-8")
            zout.writestr(item, data)
    tmp.replace(path)


def build_docx() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    if not SOURCE_DOCX.is_file():
        raise FileNotFoundError(SOURCE_DOCX)
    doc = Document(SOURCE_DOCX)
    clear_document_body(doc)
    configure_styles(doc)
    set_update_fields(doc)
    sec = doc.sections[0]
    sec.top_margin = Cm(2.5)
    sec.bottom_margin = Cm(2.5)
    sec.left_margin = Cm(2.5)
    sec.right_margin = Cm(2.5)
    sec.header_distance = Cm(1.5)
    sec.footer_distance = Cm(1.5)
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.footer.is_linked_to_previous = False
    for p in list(sec.footer.paragraphs):
        p._element.getparent().remove(p._element)
    add_page_field(sec.footer.add_paragraph())

    doc.core_properties.title = "某污水处理厂设计计算说明书（第二次修改稿）"
    doc.core_properties.subject = "环境工程课程设计"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.keywords = "污水处理厂; A2O; 课程设计"

    add_cover(doc)
    add_front_matter(doc)
    add_chapter_1(doc)
    add_chapter_2(doc)
    add_chapter_3(doc)
    add_chapter_4(doc)
    add_chapter_5(doc)
    add_chapter_6(doc)
    add_chapter_7(doc)
    add_chapter_8(doc)
    add_chapter_9(doc)
    add_chapter_10(doc)
    add_chapter_11(doc)
    add_chapter_12(doc)
    add_chapter_13(doc)
    add_chapter_14(doc)
    add_chapter_15(doc)
    add_chapter_16(doc)
    add_references(doc)

    for p in doc.paragraphs:
        p.paragraph_format.widow_control = True
        if p.style.name.startswith("Heading"):
            p.paragraph_format.keep_with_next = True
    doc.save(OUTPUT_DOCX)
    strip_comments_and_personal_metadata(OUTPUT_DOCX)


def build_xlsx() -> None:
    wb = load_workbook(SOURCE_XLSX)
    ws = wb[wb.sheetnames[0]]
    final = {
        "A1": ("已完成", "采用两格进水井、6台提升泵4用2备；单泵1687.5 m³/h、14 m、90 kW，相关正文、表5-1、设备汇总及图3-1/5-1已统一。"),
        "A2": ("已完成", "鼓风机统一为6台4用2备、单台12000 Nm³/h、65 kPa(g)；曝气器9600只，正文、设备表与图10-1一致。"),
        "A3": ("已完成", "1266/4219 m³/h明确为最高时单系列单台工作泵流量；RAS与内回流均为6台4用2备，并完成扬程复核。"),
        "A4": ("已完成", "重建19段水力高程：细格栅前至计量后5.50 m，至东厂界6.05 m；提升泵所需13.677882 m。"),
        "B1": ("已完成", "保留必要图表题名，清除红色标记；图3-1与图9-1已整图重绘。"),
        "B2": ("已完成", "保留构成完整计算链的表格项目，清除红色标记；数据按第02—08组冻结结果更新。"),
        "B3": ("已完成", "术语统一为“管道中心线长度”，坐标折线按欧氏距离逐段相加。"),
        "C1": ("已完成", "表4-1改为栅隙、净空隙数与渠道尺寸；表4-2改为栅条、槽长与水头损失。"),
        "C2": ("已处理", "图示任务书夏季东南来风；正文不再声称管理区已处于全年上风向，并列明实际工程需补全年风玫瑰。"),
        "C3": ("已复核", "主车道6 m、次通道4 m、池外砂水分离、n=0.013及进厂管2.06 m/s/2.78‰均作为课程采用值写明适用边界。"),
        "C4": ("已完成", "12幅图均从空白画布整图重绘，正文、表格、图片采用同一冻结数据；另保留可编辑SVG。"),
    }
    green = PatternFill("solid", fgColor="C6EFCE")
    for row in range(2, ws.max_row + 1):
        key = str(ws.cell(row, 1).value)
        status = ws.cell(row, 2)
        if key in final:
            status.value, ws.cell(row, 7).value = final[key]
        elif status.value == "已完成":
            status.value = "已完成"
        status.fill = green
        status.font = Font(name="宋体", size=10, color="006100")
        for col in range(1, ws.max_column + 1):
            ws.cell(row, col).alignment = Alignment(vertical="top", wrap_text=True)
            ws.cell(row, col).font = Font(name="宋体", size=10, color=ws.cell(row, col).font.color)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:I{ws.max_row}"

    if "第10组终检说明" in wb.sheetnames:
        del wb["第10组终检说明"]
    note = wb.create_sheet("第10组终检说明")
    note.append(["项目", "终检结论"])
    notes = [
        ("终检日期", "2026-09-14"),
        ("57项状态", "46项原已完成，11项由第02—09组成果闭合；终检表共57项均给出最终状态。"),
        ("流量", "Qavg=108000 m³/d；Kz=1.50；Qmax=6750 m³/h。"),
        ("提升", "两格进水井；6泵4用2备；1687.5 m³/h、14 m、90 kW。"),
        ("生物与供气", "4个A²/O系列；MLSS=3.50 kg/m³；48000 Nm³/h；9600盘；6台鼓风机4用2备。"),
        ("配水消毒", "2座二沉配水井各4支路；接触池2格；次氯酸钠5 mg/L名义、15 mg/L容量。"),
        ("水力高程", "细格栅前EL34.80；计量后EL29.30；东厂界EL28.75；落差5.50/6.05 m。"),
        ("课程未闭合项", "无。封面组号、班级和成员由提交人填写；组号洪水式需代入实际M复核。"),
        ("实际工程边界", "实测水质、最低水温、风玫瑰、厂外排口、试验及厂家曲线不属于课程阶段已具备资料。"),
    ]
    for item in notes:
        note.append(item)
    for row in note.iter_rows():
        for cell in row:
            cell.font = Font(name="宋体", size=10, bold=(cell.row == 1))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    note.column_dimensions["A"].width = 18
    note.column_dimensions["B"].width = 100
    note.freeze_panes = "A2"
    wb.save(OUTPUT_XLSX)


def text_of_doc(doc: Document) -> str:
    values = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            values.extend(cell.text for cell in row.cells)
    return "\n".join(values)


def audit() -> dict:
    doc = Document(OUTPUT_DOCX)
    text = text_of_doc(doc)
    old_patterns = {
        "四格进水井": r"四格进水井",
        "3用1备提升泵": r"3用1备",
        "旧单泵2250": r"单泵[^\n]{0,30}2250\s*m³/h",
        "旧泵132kW": r"132\s*kW",
        "旧泵H13": r"H\s*[=＝]?\s*13(?:\.0)?\s*m",
        "旧曝气器10800": r"(?<!\d)10800(?!\d)",
        "旧风机10台": r"10台(?:、|，)?9用1备|10台鼓风机",
        "旧单机6000风量": r"6000\s*Nm³/h",
        "旧总MLSS": r"4\.379",
        "旧化学固体": r"2768\.84",
        "旧SLR": r"148\.55",
        "旧甲醇容量": r"12000\s*kg/d",
        "旧铁盐容量": r"7\.47\s*m³/d",
        "旧水力总差": r"3\.90\s*m",
        "旧细格栅水位": r"EL\s*34\.30|细格栅前(?:水位为)?EL?34\.30",
        "旧控制井水位": r"EL\s*30\.40",
        "旧二沉配水": r"4座二沉配水井|每井2路",
        "过程性称谓": r"用户确认|教师意见|老师意见|Codex|待修改|黄色标记|红色标记",
    }
    old_hits = {name: len(re.findall(pattern, text)) for name, pattern in old_patterns.items()}
    required = {
        "Kz=1.50": "Kz=1.50" in text,
        "两格进水井": "两格进水井" in text,
        "提升泵6台4用2备": "6台，4用2备" in text and "1687.5 m³/h" in text,
        "A2O四系列": "4个工艺系列" in text,
        "曝气48000": "48000 Nm³/h" in text,
        "曝气器9600": "9600只" in text,
        "鼓风机6台4用2备": "设6台高速离心鼓风机、4用2备" in text,
        "二沉8座D37": "8座辐流二沉池" in text and "净径37 m" in text,
        "二沉2井4支": "二沉采用2座配水井" in text and "4条独立DN900支路" in text,
        "水力5.50/6.05": "5.50 m" in text and "6.05 m" in text,
        "东界28.75": "EL28.75 m" in text,
        "次氯酸钠容量": "15 mg/L作为设备能力上界" in text,
        "参考文献": "参考文献" in text,
    }
    with zipfile.ZipFile(OUTPUT_DOCX) as zf:
        names = zf.namelist()
        bad_parts = [n for n in names if "comments" in n.lower() or n.endswith("people.xml")]
        document_xml = zf.read("word/document.xml")
        settings_xml = zf.read("word/settings.xml")
        media = [n for n in names if n.startswith("word/media/")]
        footer_xml = b"".join(zf.read(n) for n in names if re.fullmatch(r"word/footer\d+\.xml", n))
    result = {
        "date": str(date.today()),
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "inline_shapes": len(doc.inline_shapes),
        "headings": sum(1 for p in doc.paragraphs if p.style.name.startswith("Heading")),
        "media_parts": len(media),
        "old_value_hits": old_hits,
        "required_values": required,
        "comments_or_people_parts": bad_parts,
        "tracked_change_nodes": len(re.findall(rb"<w:(?:ins|del)(?:\s|>)", document_xml)),
        "highlight_nodes": document_xml.count(b"<w:highlight"),
        "toc_field_present": b"TOC " in document_xml,
        "page_field_present": b" PAGE " in footer_xml,
        "update_fields_on_open": b"updateFields" in settings_xml,
    }
    result["passed"] = (
        not any(old_hits.values())
        and all(required.values())
        and not bad_parts
        and result["tracked_change_nodes"] == 0
        and result["highlight_nodes"] == 0
        and result["inline_shapes"] == 12
        and result["toc_field_present"]
        and result["page_field_present"]
        and result["update_fields_on_open"]
    )
    OUTPUT_JSON.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    if not result["passed"]:
        raise SystemExit(json.dumps(result, ensure_ascii=False, indent=2))
    return result


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--write", action="store_true", help="write the DOCX/XLSX outputs")
    args = parser.parse_args()
    if args.write:
        build_docx()
        build_xlsx()
    if not OUTPUT_DOCX.exists() or not OUTPUT_XLSX.exists():
        raise SystemExit("outputs do not exist; run with --write")
    result = audit()
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
