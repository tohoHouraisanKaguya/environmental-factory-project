"""Extract visible Word colour contexts and the 2026-09-13 Excel status table.

This script is read-only with respect to the common source files.  It emits a
JSON snapshot that can be used to build the group-01 location register.
"""

from __future__ import annotations

import argparse
import json
import subprocess
from pathlib import Path

from docx import Document
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
PUBLIC = ROOT / "第二次修改协作项目_2026-09-13" / "00_公共文件"
ORIGINAL = PUBLIC / "某污水处理厂设计计算说明书_原批注版.docx"
FIRST_REVISION = PUBLIC / "某污水处理厂设计计算说明书_第一版修改稿.docx"
STATUS_BOOK = PUBLIC / "修改与待修改_20260913.xlsx"


def colour_name(run) -> list[str]:
    names: list[str] = []
    highlight = run.font.highlight_color
    if highlight is not None:
        value = str(highlight).upper()
        if "YELLOW" in value or value.endswith("(7)"):
            names.append("黄色")
        elif "RED" in value or value.endswith("(6)"):
            names.append("红色高亮")
        else:
            names.append(f"其他高亮:{value}")
    rgb = run.font.color.rgb if run.font.color is not None else None
    if rgb is not None:
        value = str(rgb).upper()
        red, green, blue = (int(value[index:index + 2], 16) for index in (0, 2, 4))
        # The source document uses EE0000 rather than literal FF0000.
        if red >= 0xC0 and green <= 0x40 and blue <= 0x40:
            names.append("红色字体")
    return names


def marked_contexts(path: Path) -> dict:
    document = Document(path)
    contexts: list[dict] = []

    def append_context(location: str, paragraphs) -> None:
        marked: list[dict] = []
        full_text = "\n".join(paragraph.text for paragraph in paragraphs).strip()
        for paragraph in paragraphs:
            for run in paragraph.runs:
                colours = colour_name(run)
                if colours:
                    marked.append({"text": run.text, "colours": colours})
        if marked:
            contexts.append(
                {
                    "location": location,
                    "full_text": full_text,
                    "marked_runs": marked,
                    "colours": sorted({item for run in marked for item in run["colours"]}),
                }
            )

    for index, paragraph in enumerate(document.paragraphs, start=1):
        append_context(f"正文段落P{index}", [paragraph])
    for table_index, table in enumerate(document.tables, start=1):
        for row_index, row in enumerate(table.rows, start=1):
            for cell_index, cell in enumerate(row.cells, start=1):
                append_context(
                    f"表T{table_index}-R{row_index}-C{cell_index}",
                    cell.paragraphs,
                )

    counts: dict[str, int] = {}
    for context in contexts:
        for colour in context["colours"]:
            counts[colour] = counts.get(colour, 0) + 1
    return {
        "file": path.name,
        "paragraph_count": len(document.paragraphs),
        "table_count": len(document.tables),
        "visible_context_count": len(contexts),
        "colour_context_counts": counts,
        "contexts": contexts,
    }


def cell_colour(cell) -> str | None:
    fill = cell.fill
    if fill is None or fill.fill_type is None:
        return None
    candidates = [fill.fgColor.rgb, fill.fgColor.indexed, fill.fgColor.theme]
    for value in candidates:
        if value is not None:
            return str(value)
    return fill.fill_type


def workbook_snapshot(path: Path) -> dict:
    workbook = load_workbook(path, data_only=False, read_only=False)
    sheets = []
    for sheet in workbook.worksheets:
        rows = []
        for row_index in range(1, sheet.max_row + 1):
            values = [sheet.cell(row_index, col).value for col in range(1, sheet.max_column + 1)]
            if not any(value not in (None, "") for value in values):
                continue
            rows.append(
                {
                    "row": row_index,
                    "values": values,
                    "fills": [cell_colour(sheet.cell(row_index, col)) for col in range(1, sheet.max_column + 1)],
                }
            )
        sheets.append(
            {
                "name": sheet.title,
                "max_row": sheet.max_row,
                "max_column": sheet.max_column,
                "rows": rows,
            }
        )
    return {"file": path.name, "sheets": sheets}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()

    baseline_commit = subprocess.run(
        ["git", "rev-parse", "HEAD"],
        cwd=ROOT,
        check=True,
        capture_output=True,
        text=True,
    ).stdout.strip()
    result = {
        "baseline_commit": baseline_commit,
        "original_docx": marked_contexts(ORIGINAL),
        "first_revision_docx": marked_contexts(FIRST_REVISION),
        "status_workbook": workbook_snapshot(STATUS_BOOK),
    }
    args.out.parent.mkdir(parents=True, exist_ok=True)
    payload = json.dumps(result, ensure_ascii=False, indent=2) + "\n"
    args.out.write_bytes(payload.encode("utf-8"))

    print(json.dumps(
        {
            "original_contexts": result["original_docx"]["visible_context_count"],
            "original_colours": result["original_docx"]["colour_context_counts"],
            "first_revision_contexts": result["first_revision_docx"]["visible_context_count"],
            "first_revision_colours": result["first_revision_docx"]["colour_context_counts"],
            "sheets": [sheet["name"] for sheet in result["status_workbook"]["sheets"]],
        },
        ensure_ascii=False,
        indent=2,
    ))


if __name__ == "__main__":
    main()
